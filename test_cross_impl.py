"""Cross-implementation checks: the HTML tab (index.html) must agree with ashtakavarga.py.

    python3 -m unittest -v test_cross_impl

These tests run the code that actually ships in index.html (extracted from between the
ASHTAKAVARGA-START / ASHTAKAVARGA-END markers) under node, and compare it with the
Python implementation.  They skip, loudly, when node is not installed.
"""
import json
import os
import pathlib
import random
import re
import shutil
import subprocess
import tempfile
import unittest

import ashtakavarga as av

ROOT = pathlib.Path(__file__).resolve().parent
INDEX = ROOT / "index.html"
NODE = shutil.which("node")

NODE_HARNESS = r"""
const fs = require('fs'), vm = require('vm');
const html = fs.readFileSync(process.argv[1], 'utf8');
const m = html.match(/\/\* ASHTAKAVARGA-START[^\n]*\n([\s\S]*?)\/\* ASHTAKAVARGA-END \*\//);
if (!m) { console.log(JSON.stringify({error: 'markers not found'})); process.exit(0); }
const ctx = vm.createContext({console});
vm.runInContext(m[1], ctx);
const job = JSON.parse(fs.readFileSync(0, 'utf8'));
const out = {};
out.rules = vm.runInContext('ASHTAKAVARGA_RULES', ctx);
out.results = job.charts.map(c => {
  ctx.__c = c;
  return vm.runInContext('computeAshtakavarga(__c)', ctx);
});
out.errors = job.bad.map(b => {
  ctx.__b = b;
  try { vm.runInContext('computeAshtakavarga(__b.chart, __b.rules || undefined)', ctx); return null; }
  catch (e) { return String(e.message); }
});
console.log(JSON.stringify(out));
"""


def run_html_block(charts, bad=()):
    p = subprocess.run([NODE, "-e", NODE_HARNESS, str(INDEX)],
                       input=json.dumps({"charts": charts, "bad": list(bad)}),
                       capture_output=True, text=True, timeout=60)
    if p.returncode != 0:
        raise AssertionError(f"node failed: {p.stderr}")
    out = json.loads(p.stdout)
    if "error" in out:
        raise AssertionError(out["error"])
    return out


def random_charts(n, seed=1008):
    rng = random.Random(seed)
    charts = [{c: rng.randrange(12) for c in av.CONTRIBUTORS} for _ in range(n)]
    charts += [{c: s for c in av.CONTRIBUTORS} for s in range(12)]  # every sign, all together
    return charts


@unittest.skipUnless(NODE, "node is not installed — HTML cross-check skipped")
class HtmlMatchesPython(unittest.TestCase):
    def test_markers_exist_exactly_once(self):
        text = INDEX.read_text(encoding="utf-8")
        self.assertEqual(text.count("ASHTAKAVARGA-START"), 1)
        self.assertEqual(text.count("ASHTAKAVARGA-END"), 1)

    def test_html_rule_table_is_identical_to_python(self):
        html_rules = run_html_block([])["rules"]
        py_rules = {t: {c: list(h) for c, h in row.items()} for t, row in av.RULES.items()}
        self.assertEqual(html_rules, py_rules)

    def test_html_compute_equals_python_on_random_charts(self):
        charts = random_charts(2000)
        got = run_html_block(charts)["results"]
        for chart, res in zip(charts, got):
            self.assertEqual(res, av.compute(chart), chart)

    def test_html_refuses_a_rule_table_that_fails_its_checksum(self):
        rules = {t: {c: list(h) for c, h in row.items()} for t, row in av.RULES.items()}
        rules["Moon"]["Moon"].append(12)  # one bindu too many
        chart = {c: 0 for c in av.CONTRIBUTORS}
        errors = run_html_block([], bad=[{"chart": chart, "rules": rules}])["errors"]
        self.assertRegex(errors[0] or "", r"Moon.*50.*49")

    def test_html_rejects_bad_signs(self):
        good = {c: 0 for c in av.CONTRIBUTORS}
        bads = []
        for label, mutate in (("sign 12", lambda c: c.update(Sun=12)),
                              ("negative", lambda c: c.update(Sun=-1)),
                              ("fraction", lambda c: c.update(Sun=1.5)),
                              ("string", lambda c: c.update(Sun="Aries")),
                              ("missing Lagna", lambda c: c.pop("Lagna"))):
            c = dict(good)
            mutate(c)
            bads.append({"chart": c, "rules": None})
        errors = run_html_block([], bad=bads)["errors"]
        for e, b in zip(errors, bads):
            self.assertTrue(e, f"expected an error for {b['chart']}")


class HtmlTabIsWired(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.html = INDEX.read_text(encoding="utf-8")

    def test_tab_button_follows_muhurta_in_the_tab_bar(self):
        bar = re.search(r'id="reportTabs">(.*?)</div>', self.html, re.S).group(1)
        tabs = re.findall(r'data-tab="([a-z]+)"', bar)
        self.assertEqual(tabs[-2:], ["muhurta", "ashtakavarga"])
        self.assertRegex(bar, r'<button type="button" class="rtab" data-tab="ashtakavarga">[^<]*Ashtakavarga[^<]*</button>')

    def test_pane_exists_hidden_with_a_print_title(self):
        self.assertRegex(
            self.html,
            r'<div class="rpane" id="pane-ashtakavarga" data-print-title="[^"]+" style="display:none"></div>')

    def test_tab_switcher_knows_the_new_pane(self):
        sw = re.search(r'\[("transit".*?)\]\.forEach\(k=>', self.html, re.S).group(1)
        self.assertIn('"ashtakavarga"', sw)

    def test_predictive_render_calls_the_new_renderer(self):
        body = re.search(r"function renderPredictive\(\)\{(.*?)\n\}", self.html, re.S).group(1)
        self.assertIn("renderAshtakavargaPane(", body)

    def test_section_subtitle_lists_the_tab(self):
        m = re.search(r'Predictive &amp; Remedies</span>\s*<span class="section-sub">([^<]*)</span>',
                      self.html)
        self.assertIsNotNone(m)
        self.assertIn("Ashtakavarga", m.group(1))


BROWSER_SCRIPT = r"""
const { chromium } = require(process.env.PW_MODULE);
const indexUrl = 'file://' + process.argv[1];
const shots = process.env.AV_ARTIFACT_DIR || '';
const BIRTHS = [
  {dob:'1990-05-15', tob:'06:30:00', lat:'19.08',  lon:'72.88',  tz:'5.5'},
  {dob:'1975-11-02', tob:'22:10:00', lat:'13.0',   lon:'77.6',   tz:'5.5'},
  {dob:'2001-02-28', tob:'13:45:00', lat:'51.5',   lon:'-0.12',  tz:'0'},
];
(async () => {
  const browser = await chromium.launch({headless: true});
  const ctx = await browser.newContext({viewport: {width: 1280, height: 900}});
  const blocked = [];
  // The page POSTs every generated chart to a Google Sheet: allow file:// only.
  await ctx.route('**/*', r => r.request().url().startsWith('file:') ? r.continue()
                                : (blocked.push(r.request().url()), r.abort()));
  const page = await ctx.newPage();
  const jsErrors = [];
  page.on('pageerror', e => jsErrors.push(String(e)));
  await page.goto(indexUrl);
  await page.click('#btn-manual');
  const out = {charts: [], jsErrors, blocked: null};
  for (const [i, b] of BIRTHS.entries()) {
    await page.fill('#f-dob', b.dob);
    await page.fill('#f-tob', b.tob);
    await page.fill('#f-lat', b.lat); await page.press('#f-lat', 'Tab');
    await page.fill('#f-lon', b.lon); await page.press('#f-lon', 'Tab');
    await page.fill('#f-tz', b.tz);
    await page.click('#btn-generate');
    await page.waitForSelector('#report.show');
    await page.click('#reportNav button:has-text("Predictive")');
    const before = await page.evaluate(() => getComputedStyle(document.getElementById('pane-ashtakavarga')).display);
    await page.click('#reportTabs .rtab[data-tab="ashtakavarga"]');
    const d = await page.evaluate(() => {
      const pane = document.getElementById('pane-ashtakavarga');
      const rows = t => [...t.querySelectorAll('tbody tr')].map(tr => [...tr.children].map(td => td.textContent.trim()));
      const cls = t => [...t.querySelectorAll('tbody tr')].map(tr => [...tr.children].map(td => td.className));
      const ts = pane.querySelectorAll('table');
      const panes = {};
      document.querySelectorAll('#report .rpane').forEach(p => panes[p.id] = getComputedStyle(p).display);
      return {
        tabs: [...document.querySelectorAll('#reportTabs .rtab')].map(b => b.dataset.tab),
        active: [...document.querySelectorAll('#reportTabs .rtab.active')].map(b => b.dataset.tab),
        panes, nTables: ts.length,
        bav: rows(ts[0]), bavCls: cls(ts[0]), sav: rows(ts[1]), savCls: cls(ts[1]),
        headers: [...ts[0].querySelectorAll('thead th')].map(th => th.innerText.replace(/\s+/g, ' ').trim()),
        sid: currentChart.chart.sidereal, lagna: currentChart.lagnaRashi.n,
        text: pane.innerText,
      };
    });
    d.paneBeforeClick = before;
    out.charts.push(d);
    if (i === 0) {
      // error states, rendered by the page itself
      out.noLagna = await page.evaluate(() => { renderAshtakavargaPane(currentChart.chart.sidereal, 0);
        const p = document.getElementById('pane-ashtakavarga'); return {text: p.innerText, tables: p.querySelectorAll('table').length}; });
      out.badRules = await page.evaluate(() => {
        ASHTAKAVARGA_RULES.Moon.Moon.push(12);
        renderAshtakavargaPane(currentChart.chart.sidereal, currentChart.lagnaRashi.n);
        const p = document.getElementById('pane-ashtakavarga'); const r = {text: p.innerText, tables: p.querySelectorAll('table').length};
        ASHTAKAVARGA_RULES.Moon.Moon.pop();
        renderAshtakavargaPane(currentChart.chart.sidereal, currentChart.lagnaRashi.n);
        r.restoredTables = p.querySelectorAll('table').length;
        return r; });
      // layout: desktop, phone, print
      out.layout = {};
      const pane = page.locator('#pane-ashtakavarga');
      if (shots) await pane.screenshot({path: shots + '/av_desktop.png'});
      await page.setViewportSize({width: 375, height: 800});
      out.layout.phone = await page.evaluate(() => {
        const vw = window.innerWidth;
        return {vw, pageScrollW: document.documentElement.scrollWidth,
                wrappers: [...document.querySelectorAll('#pane-ashtakavarga .table-scroll')].map(w => {
                  const tb = w.querySelector('table');
                  return {w: Math.round(w.getBoundingClientRect().width),
                          scrolls: w.scrollWidth > w.clientWidth || tb.scrollWidth > tb.clientWidth}; })};
      });
      if (shots) await pane.screenshot({path: shots + '/av_phone.png'});
      await page.setViewportSize({width: 1280, height: 900});
      await page.emulateMedia({media: 'print'});
      out.layout.print = await page.evaluate(() => {
        const pane = document.getElementById('pane-ashtakavarga');
        const strong = pane.querySelector('td.av-strong'), weak = pane.querySelector('td.av-weak');
        return {display: getComputedStyle(pane).display,
                strongBg: strong && getComputedStyle(strong).backgroundColor,
                weakBg: weak && getComputedStyle(weak).backgroundColor,
                title: getComputedStyle(pane, '::before').content,
                exact: strong && (getComputedStyle(strong).printColorAdjust || getComputedStyle(strong).webkitPrintColorAdjust)};
      });
      // A4 less the site's 13 mm margins is 184 mm = ~695 px, just above the 680 px phone breakpoint.
      await page.setViewportSize({width: 695, height: 1000});
      out.layout.printA4 = await page.evaluate(() =>
        [...document.querySelectorAll('#pane-ashtakavarga .table-scroll')].map(w => {
          const t = w.querySelector('table');
          return {wrapClient: w.clientWidth, wrapScroll: w.scrollWidth, tableScroll: t.scrollWidth,
                  tableWidth: Math.round(t.getBoundingClientRect().width),
                  headers: [...t.querySelectorAll('thead th')].map(th => ({txt: th.innerText.replace(/\s+/g, ' ').trim(), cw: th.clientWidth,
                                                              sw: (() => { const r = document.createRange(); r.selectNodeContents(th); return Math.ceil(r.getBoundingClientRect().width); })()}))};
        }));
      await page.setViewportSize({width: 1280, height: 900});
      if (shots) await page.pdf({path: shots + '/av_print.pdf', format: 'A4', printBackground: true});
      await page.emulateMedia({media: 'screen'});
    }
  }
  out.blocked = blocked;
  console.log(JSON.stringify(out));
  await browser.close();
})().catch(e => { console.error(e); process.exit(1); });
"""


def _playwright_module():
    npm = shutil.which("npm")
    if not npm:
        return None
    root = subprocess.run([npm, "root", "-g"], capture_output=True, text=True).stdout.strip()
    cand = pathlib.Path(root) / "playwright"
    return str(cand) if cand.exists() else None


PW_MODULE = _playwright_module() if NODE else None


@unittest.skipUnless(PW_MODULE and os.path.isdir(os.environ.get("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")),
                     "playwright + chromium not available — browser test skipped")
class BrowserEndToEnd(unittest.TestCase):
    """Generate real horoscopes in headless Chromium and read the rendered tab."""

    @classmethod
    def setUpClass(cls):
        env = dict(os.environ, PW_MODULE=PW_MODULE)
        env.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
        p = subprocess.run([NODE, "-e", BROWSER_SCRIPT, str(INDEX)], capture_output=True,
                           text=True, timeout=180, env=env)
        if p.returncode != 0:
            raise AssertionError(f"browser run failed:\n{p.stderr[-2000:]}")
        cls.out = json.loads(p.stdout.strip().splitlines()[-1])

    def test_no_javascript_errors_and_no_network_escaped(self):
        self.assertEqual(self.out["jsErrors"], [])
        # everything that tried to leave the machine (fonts, the sheet logger) was blocked
        self.assertTrue(all(not u.startswith("file:") for u in self.out["blocked"]))

    def test_tab_is_last_and_switches_panes(self):
        for d in self.out["charts"]:
            self.assertEqual(d["tabs"][-2:], ["muhurta", "ashtakavarga"])
            self.assertEqual(d["active"], ["ashtakavarga"])
            self.assertEqual({k for k, v in d["panes"].items() if v != "none"}, {"pane-ashtakavarga"})
        self.assertEqual(self.out["charts"][0]["paneBeforeClick"], "none")

    def test_rendered_numbers_equal_python_for_each_chart(self):
        for d in self.out["charts"]:
            signs = {p: int(d["sid"][p] % 360 // 30) for p in av.PLANETS}
            signs["Lagna"] = d["lagna"] - 1
            want = av.compute(signs)
            self.assertEqual(d["nTables"], 2)
            self.assertEqual(len(d["bav"]), 7)
            for planet, row in zip(av.PLANETS, d["bav"]):
                self.assertTrue(row[0].startswith(planet), row[0])
                self.assertIn(av.SIGNS[signs[planet]], row[0])  # "Sun in Aquarius"
                self.assertEqual([int(x) for x in row[1:13]], want["bav"][planet], planet)
                self.assertEqual(int(row[13].split()[0]), av.EXPECTED_TOTALS[planet])
                self.assertIn("✓", row[13])
            sav_row = d["sav"][1]
            self.assertEqual([int(x) for x in sav_row[1:13]], want["sav"])
            self.assertEqual(int(sav_row[13].split()[0]), 337)
            house_row = d["sav"][0]
            self.assertEqual([int(x) for x in house_row[1:13]],
                             [(s - signs["Lagna"]) % 12 + 1 for s in range(12)])

    def test_charts_differ_so_the_pane_is_not_stale(self):
        savs = [tuple(int(x) for x in d["sav"][1][1:13]) for d in self.out["charts"]]
        self.assertEqual(len(set(savs)), 3, savs)

    def test_shading_and_natal_sign_outline(self):
        for d in self.out["charts"]:
            signs = {p: int(d["sid"][p] % 360 // 30) for p in av.PLANETS}
            for planet, row, cls in zip(av.PLANETS, d["bav"], d["bavCls"]):
                for i, (v, c) in enumerate(zip(row[1:13], cls[1:13])):
                    v = int(v)
                    band = "av-strong" if v >= 5 else "av-avg" if v >= 3 else "av-weak"
                    self.assertIn(band, c.split(), (planet, i, v, c))
                    self.assertEqual("av-self" in c.split(), i == signs[planet], (planet, i))
            for v, c in zip(d["sav"][1][1:13], d["savCls"][1][1:13]):
                v = int(v)
                band = "av-strong" if v >= 30 else "av-avg" if v >= 25 else "av-weak"
                self.assertIn(band, c.split(), (v, c))

    def test_pane_states_the_source_of_the_tables(self):
        for d in self.out["charts"]:
            self.assertIn("Bṛhat Parāśara Horā Śāstra, ch. 66", d["text"])

    def test_lagna_sign_is_flagged_in_both_headers(self):
        for d in self.out["charts"]:
            self.assertEqual(sum("Lagna" in h for h in d["headers"]), 1)

    def test_missing_lagna_shows_a_message_not_wrong_numbers(self):
        self.assertEqual(self.out["noLagna"]["tables"], 0)
        self.assertIn("needs the Lagna", self.out["noLagna"]["text"])

    def test_a_corrupt_rule_table_stops_the_calculation_in_the_page(self):
        bad = self.out["badRules"]
        self.assertEqual(bad["tables"], 0)
        self.assertRegex(bad["text"], r"Calculation stopped.*Moon: table has 50 bindus")
        self.assertEqual(bad["restoredTables"], 2)

    def test_phone_tables_scroll_inside_their_own_box(self):
        ph = self.out["layout"]["phone"]
        for w in ph["wrappers"]:
            self.assertLessEqual(w["w"], ph["vw"], ph)
        # 14 columns cannot fit a phone: each table must scroll inside its own box
        # (the site's mobile CSS makes .rtable itself the scroll container).
        self.assertEqual(len(ph["wrappers"]), 2)
        self.assertTrue(all(w["scrolls"] for w in ph["wrappers"]), ph)

    def test_printed_tables_fit_the_a4_page_with_all_twelve_signs_and_the_total(self):
        wraps = self.out["layout"]["printA4"]
        self.assertEqual(len(wraps), 2)
        for w in wraps:
            self.assertLessEqual(w["wrapScroll"], w["wrapClient"] + 1, w)   # nothing clipped
            self.assertLessEqual(w["tableScroll"], w["wrapClient"] + 1, w)
            texts = [h["txt"].lower() for h in w["headers"]]  # print CSS shows labels in capitals
            joined = " ".join(texts)
            for sign in av.SIGNS:  # three-letter abbreviations on paper
                self.assertIn(sign[:3].lower(), joined, (sign, texts))
            self.assertIn("total", texts)
            for h in w["headers"]:  # no header text may spill into its neighbour
                self.assertLessEqual(h["sw"], h["cw"] + 1, h)

    def test_print_keeps_the_pane_title_and_shading(self):
        pr = self.out["layout"]["print"]
        self.assertNotEqual(pr["display"], "none")
        self.assertIn("Ashtakavarga", pr["title"])
        self.assertNotIn(pr["strongBg"], (None, "rgba(0, 0, 0, 0)"))
        self.assertNotIn(pr["weakBg"], (None, "rgba(0, 0, 0, 0)"))
        self.assertEqual(pr["exact"], "exact")


SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

SHIPPED_XLSX = ROOT / "Ashtakavarga.xlsx"
SAMPLE_CHART = {"Sun": "Taurus", "Moon": "Sagittarius", "Mars": "Aquarius", "Mercury": "Aries",
                "Jupiter": "Gemini", "Venus": "Pisces", "Saturn": "Capricorn", "Lagna": "Taurus"}
BAV_ROW0, SAV_ROW, INPUT_ROW0 = 16, 27, 5  # layout contract of the sheet "Ashtakavarga"


def recalc_with_libreoffice(paths, outdir):
    profile = tempfile.mkdtemp(prefix="lo_profile_")
    try:
        subprocess.run([SOFFICE, f"-env:UserInstallation=file://{profile}", "--headless",
                        "--convert-to", "xlsx", "--outdir", str(outdir), *map(str, paths)],
                       check=True, capture_output=True, timeout=300)
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def read_results(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Ashtakavarga"]
    bav = {p: [ws.cell(row=BAV_ROW0 + i, column=2 + s).value for s in range(12)]
           for i, p in enumerate(av.PLANETS)}
    sav = [ws.cell(row=SAV_ROW, column=2 + s).value for s in range(12)]
    checks = [ws.cell(row=BAV_ROW0 + i, column=16).value for i in range(7)]
    return {"bav": bav, "sav": sav}, checks, ws.cell(row=SAV_ROW, column=16).value


@unittest.skipUnless(SOFFICE and openpyxl, "LibreOffice and openpyxl are needed for the Excel check")
class ExcelMatchesPython(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import make_ashtakavarga_xlsx as mx
        cls.mx = mx
        cls.tmp = pathlib.Path(tempfile.mkdtemp(prefix="av_xlsx_"))
        cls.addClassCleanup(shutil.rmtree, cls.tmp, ignore_errors=True)
        cls.base = cls.tmp / "base.xlsx"
        mx.build(cls.base)
        rng = random.Random(77)
        cls.charts = [{c: s for c in av.CONTRIBUTORS} for s in range(12)]
        cls.charts += [{c: rng.randrange(12) for c in av.CONTRIBUTORS} for _ in range(12)]
        inputs = []
        for i, chart in enumerate(cls.charts):
            wb = openpyxl.load_workbook(cls.base)
            ws = wb["Ashtakavarga"]
            for k, c in enumerate(av.CONTRIBUTORS):
                ws.cell(row=INPUT_ROW0 + k, column=2).value = av.SIGNS[chart[c]]
            f = cls.tmp / f"chart_{i:02d}.xlsx"
            wb.save(f)
            inputs.append(f)
        wb = openpyxl.load_workbook(cls.base)  # one deliberately corrupted rule table
        wb["Rules"].cell(row=mx.rules_row("Moon", "Moon"), column=2 + 12).value = 1
        cls.corrupt = cls.tmp / "corrupt.xlsx"
        wb.save(cls.corrupt)
        cls.out = cls.tmp / "recalc"
        cls.out.mkdir()
        recalc_with_libreoffice(inputs + [cls.corrupt, cls.base], cls.out)

    def test_every_bav_and_sav_cell_equals_python_on_24_charts(self):
        for i, chart in enumerate(self.charts):
            got, _, _ = read_results(self.out / f"chart_{i:02d}.xlsx")
            self.assertEqual(got, av.compute(chart), chart)

    def test_default_sample_chart_matches_python(self):
        signs = {c: av.SIGNS.index(n) for c, n in SAMPLE_CHART.items()}
        got, checks, sav_check = read_results(self.out / "base.xlsx")
        self.assertEqual(got, av.compute(signs))
        self.assertEqual(checks, ["OK"] * 7)
        self.assertEqual(sav_check, "OK")

    def test_a_corrupted_rule_table_is_flagged_not_silently_used(self):
        _, checks, sav_check = read_results(self.out / "corrupt.xlsx")
        self.assertEqual(checks[av.PLANETS.index("Moon")], "CHECK")
        self.assertEqual(sav_check, "CHECK")
        self.assertEqual(checks.count("OK"), 6)

    def test_results_are_live_formulas_and_inputs_are_plain_values(self):
        wb = openpyxl.load_workbook(self.base)
        ws, work = wb["Ashtakavarga"], wb["Workings"]
        for i in range(7):
            for s in range(12):
                self.assertTrue(str(ws.cell(row=BAV_ROW0 + i, column=2 + s).value).startswith("="))
        for s in range(12):
            self.assertTrue(str(ws.cell(row=SAV_ROW, column=2 + s).value).startswith("="))
        for k, c in enumerate(av.CONTRIBUTORS):
            self.assertEqual(ws.cell(row=INPUT_ROW0 + k, column=2).value, SAMPLE_CHART[c])
            self.assertEqual(ws.cell(row=INPUT_ROW0 + k, column=1).value, c)
        self.assertGreater(sum(str(c.value).startswith("=") for row in work.iter_rows() for c in row), 600)

    def test_rules_sheet_holds_exactly_the_python_tables(self):
        ws = openpyxl.load_workbook(self.base)["Rules"]
        for t in av.PLANETS:
            for c in av.CONTRIBUTORS:
                r = self.mx.rules_row(t, c)
                self.assertEqual(ws.cell(row=r, column=1).value, t)
                houses = [h for h in range(1, 13) if ws.cell(row=r, column=2 + h).value == 1]
                self.assertEqual(houses, sorted(av.RULES[t][c]), (t, c))
                self.assertEqual(ws.cell(row=r, column=2).value, c)

    def test_row_labels_are_intact_and_no_heading_overwrites_a_table_row(self):
        ws = openpyxl.load_workbook(self.out / "base.xlsx", data_only=True)["Ashtakavarga"]
        self.assertEqual([ws.cell(row=BAV_ROW0 + i, column=1).value for i in range(7)], list(av.PLANETS))
        self.assertEqual(ws.cell(row=SAV_ROW, column=1).value, "SAV bindus")
        self.assertEqual(ws.cell(row=SAV_ROW - 1, column=1).value, "House from Lagna")
        self.assertEqual([ws.cell(row=BAV_ROW0 - 1, column=2 + s).value for s in range(12)], list(av.SIGNS))
        self.assertEqual([ws.cell(row=SAV_ROW - 2, column=2 + s).value for s in range(12)], list(av.SIGNS))
        # section headings sit on their own rows, clear of both tables
        headings = [r for r in range(1, 40) if str(ws.cell(row=r, column=1).value or "").startswith(
            ("Bhinnashtakavarga", "Sarvashtakavarga"))]
        self.assertEqual(len(headings), 2)
        for r in headings:
            self.assertIsNone(ws.cell(row=r, column=2).value, r)
            self.assertNotIn(ws.cell(row=r, column=1).value, list(av.PLANETS))
        self.assertEqual(headings[1], SAV_ROW - 3)

    def test_rules_sheet_names_the_source_and_the_moon_variant(self):
        ws = openpyxl.load_workbook(self.base)["Rules"]
        note = " ".join(str(ws.cell(row=r, column=1).value or "") for r in (2, 3))
        self.assertIn("ch. 66", note)
        self.assertIn("Moon", note)
        self.assertIn("remove 9", note)

    def test_sign_inputs_have_a_dropdown_of_the_twelve_signs(self):
        ws = openpyxl.load_workbook(self.base)["Ashtakavarga"]
        dvs = [dv for dv in ws.data_validations.dataValidation if dv.type == "list"]
        self.assertTrue(any("B5" in str(dv.sqref) for dv in dvs))


@unittest.skipUnless(openpyxl, "openpyxl is needed")
class ShippedWorkbook(unittest.TestCase):
    def test_committed_xlsx_exists_with_formulas_and_cached_values(self):
        self.assertTrue(SHIPPED_XLSX.exists(), "run: python3 make_ashtakavarga_xlsx.py")
        live = openpyxl.load_workbook(SHIPPED_XLSX)["Ashtakavarga"]
        cached, _, sav_check = read_results(SHIPPED_XLSX)
        self.assertTrue(str(live.cell(row=BAV_ROW0, column=2).value).startswith("="), "formulas must survive")
        signs = {c: av.SIGNS.index(live.cell(row=INPUT_ROW0 + k, column=2).value)
                 for k, c in enumerate(av.CONTRIBUTORS)}
        self.assertEqual(cached, av.compute(signs), "cached values must match Python for the saved inputs")
        self.assertEqual(sav_check, "OK")


if __name__ == "__main__":
    unittest.main()
