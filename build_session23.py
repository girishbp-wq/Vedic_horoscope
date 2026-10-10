#!/usr/bin/env python3
"""Build the Session 23 rules from the master workbook:

    Classification_for_Horoscope_Analysis_v7_1.xlsx  (sheets prefixed S23_)
        ->  session23_rules.json  +  the SESSION23-DATA block of index.html

    python3 build_session23.py                  # master workbook found next to this file
    python3 build_session23.py --workbook FILE  # any workbook that has the S23_ sheets (e.g. an --export copy)
    python3 build_session23.py --check          # only check the S23_ sheets (publish.bat runs this first)

`publish.bat` runs this right after scripts/build_data.py, so one double-click rebuilds everything.
The S23_ sheets are the editable source of truth (edit a text or a class rule, save, publish).
`reference` in the JSON is a copy of the page's own tables (bhāva significations, rāśi classifications,
dignity, aspects, daśā constants).  It is read from the master workbook with the same loaders
scripts/build_data.py uses; where that script is not available it is read from index.html through node.
"""
import importlib.util
import json
import pathlib
import shutil
import subprocess
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parent
JSON_PATH = ROOT / "session23_rules.json"
INDEX = ROOT / "index.html"
MASTER_GLOB = "Classification_for_Horoscope_Analysis_*.xlsx"
PREFIX = "S23_"

HEAD_FILL = PatternFill("solid", fgColor="6B1D2B")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")

PROFILE_KEYS = ["Relationships", "Profession / Status", "Reputation", "Personality Traits",
                "Body Parts Ruled", "General Health Indicator", "Nature", "Combustion Effect",
                "Soul / Mind / Body"]
REFERENCE_NAMES = ["BHAVA_INFO", "RASHI", "RASHI_DIGNITY", "DIGNITY_DEG", "PLANET_OWN_HOUSES",
                   "NATURAL_FRIENDS", "KARAKATWAS", "SPECIAL_ASPECTS", "VORDER", "VYEARS",
                   "COMBUST_ORB", "COMBUST_PLANETS", "NAKSHATRAS", "PLANET_PROFILE"]
# Constants the page holds itself (not in the workbook); test_session23_data checks them against the live page.
FIXED_REFERENCE = {
    "VORDER": ["Ketu", "Venus", "Sun", "Moon", "Mars", "Rahu", "Jupiter", "Saturn", "Mercury"],
    "VYEARS": {"Ketu": 7, "Venus": 20, "Sun": 6, "Moon": 10, "Mars": 7, "Rahu": 18, "Jupiter": 16, "Saturn": 19, "Mercury": 17},
    "COMBUST_ORB": 5,
    "COMBUST_PLANETS": ["Moon", "Mars", "Mercury", "Jupiter", "Venus", "Saturn"],
}

# sheet name -> header row.  (The data sheets only; calculator sheets are added by make_session23_xlsx.py.)
SHEETS = {
    PREFIX + "Classes": ["No", "Class", "SlideName", "Houses", "OtherNames", "Nature", "Rule", "Slides"],
    PREFIX + "ClassRules": ["Class", "AppliesTo", "ExcludeHouses", "Text", "Slide"],
    PREFIX + "Badhaka": ["Mode", "House"],
    PREFIX + "DignityEffect": ["Label", "Band", "Text", "Status", "Slide"],
    PREFIX + "Digbala": ["Planet", "StrongHouse", "LostHouse", "Status", "Source"],
    PREFIX + "DashaRoleText": ["Role", "Kind", "Text", "Slide"],
    PREFIX + "GrahaInBhava": ["Planet", "House", "Points", "Extra", "Status", "Source"],
    PREFIX + "GrahaPair": ["A", "B", "Conjunction", "Aspect", "Status", "Source"],
    PREFIX + "BhavaLordIn": ["LordOf", "SitsIn", "Text", "Status", "Source", "Condition", "Exchange"],
    PREFIX + "GrahaRashi": ["Planet", "Rashi", "Text", "Status", "Source"],
    PREFIX + "BhavaNature": ["House", "Nature", "Planet", "Text", "Status", "Source"],
    PREFIX + "AspectMeaning": ["Planet", "Aspect", "FromHouse", "Text", "Status", "Source"],
    PREFIX + "LifeAreas": ["No", "Area", "Houses", "Karaka", "KarakaFemale", "Link", "Source"],
    PREFIX + "Conditions": ["Key", "Planet", "House", "Text", "Status", "Source"],
    PREFIX + "Remedies": ["Topic", "Text", "Source"],
}
LONG_COLUMNS = ("Nature", "Rule", "Text", "Points", "Extra", "Conjunction", "Aspect", "Source", "OtherNames")

README_LINES = [
    "Session 23 rules — editable source of truth (sheets whose names start with S23_)",
    "Edit any cell in the S23_ data sheets, save, then double-click publish.bat (or run python build_session23.py).",
    "Status column: taught = the teacher's own content (slides / recording);  curated = drafted by blending her karakatwas, not yet taught;  blend = generated from templates;  standard = a standard rule the teacher has not stated yet.",
    "Houses are comma lists (1, 4, 7, 10). Points: one paragraph per blank-line-separated block; Extra: one line each.",
    "Keep the S23_ClassRules rows grouped in the order of the ten classes (Kendra … Trishadaya): the Excel calculator joins their texts in sheet order.",
    "S23_Chart, S23_Classes_Calc, S23_Roles_Calc, S23_Dasha_Calc, S23_Predict_Calc and S23_Ref_Calc are the live calculator (formulas) — set the signs on S23_Chart.",
    "When a later session teaches a graha, replace its 'curated' rows and set Status to taught.",
    "Sessions 23 (2024 deck) to 27: S23_BhavaNature (benefic/malefic in each house), S23_AspectMeaning, S23_LifeAreas "
    "(Ready Reckoner), S23_Conditions (sentences that hold only for some charts; Key is one of a fixed list) and S23_Remedies. "
    "Blank House / Planet / Aspect / FromHouse means 'any'. In S23_BhavaLordIn, Condition is strong / weak / blank and "
    "Exchange is yes for a Parivartana row.",
]


# ---------------------------------------------------------------- sheets <-> rules dict
def _ints(text):
    return [int(x) for x in str(text).replace(";", ",").split(",") if str(x).strip()]


def _join_ints(values):
    return ", ".join(str(v) for v in values)


def rules_to_rows(rules):
    """The rules dict as sheet rows, keyed by (prefixed) sheet name — the inverse of read_workbook()."""
    return {
        PREFIX + "Classes": [[c["no"], c["name"], c["slide_name"], _join_ints(c["houses"]) if c["houses"] else "by Lagna",
                              c["other_names"], c["nature"], c["rule"], c["slides"]] for c in rules["classes"]],
        PREFIX + "ClassRules": [[r["class"], r["applies"], _join_ints(r["exclude_houses"]), r["text"], r["slide"]]
                                for r in rules["class_rules"]],
        PREFIX + "Badhaka": [[m, h] for m, h in rules["badhaka"].items()],
        PREFIX + "DignityEffect": [[k, v["band"], v["text"], v["status"], v["slide"]] for k, v in rules["dignity_effect"].items()],
        PREFIX + "Digbala": [[p, v["strong"], v["lost"], v["status"], v["source"]] for p, v in rules["digbala"].items()],
        PREFIX + "DashaRoleText": [[k, v["kind"], v["text"], v["slide"]] for k, v in rules["dasha_role_text"].items()],
        PREFIX + "GrahaInBhava": [[r["planet"], r["house"], "\n\n".join(r["points"]), "\n".join(r["extra"]), r["status"], r["source"]]
                                  for r in rules["graha_in_bhava"]],
        PREFIX + "GrahaPair": [[r["a"], r["b"], r["conjunction"], r["aspect"], r["status"], r["source"]] for r in rules["graha_pair"]],
        PREFIX + "BhavaLordIn": [[r["lord_of"], r["sits_in"], r["text"], r["status"], r["source"], r["condition"],
                                  "yes" if r["exchange"] else ""] for r in rules["bhava_lord_in"]],
        PREFIX + "GrahaRashi": [[r["planet"], r["rashi"], r["text"], r["status"], r["source"]] for r in rules["graha_rashi"]],
        PREFIX + "BhavaNature": [[r["house"], r["nature"], r["planet"], r["text"], r["status"], r["source"]] for r in rules["bhava_nature"]],
        PREFIX + "AspectMeaning": [[r["planet"], r["aspect"], r["from_house"], r["text"], r["status"], r["source"]]
                                   for r in rules["aspect_meaning"]],
        PREFIX + "LifeAreas": [[r["no"], r["area"], _join_ints(r["houses"]), ", ".join(r["karakas"]), r["karaka_female"], r["link"],
                                r["source"]] for r in rules["life_areas"]],
        PREFIX + "Conditions": [[r["key"], r["planet"], r["house"], r["text"], r["status"], r["source"]] for r in rules["conditions"]],
        PREFIX + "Remedies": [[r["topic"], r["text"], r["source"]] for r in rules["remedies"]],
    }


def write_data_sheets(wb, rules):
    """Create (replacing any existing) the S23_ data sheets in `wb` from the rules dict."""
    rows = rules_to_rows(rules)
    for name, header in SHEETS.items():
        if name in wb.sheetnames:
            del wb[name]
        sh = wb.create_sheet(name)
        for j, h in enumerate(header, 1):
            c = sh.cell(row=1, column=j, value=h)
            c.fill, c.font = HEAD_FILL, HEAD_FONT
        for i, row in enumerate(rows[name], 2):
            for j, v in enumerate(row, 1):
                sh.cell(row=i, column=j, value=v).alignment = WRAP
        for j, h in enumerate(header, 1):
            sh.column_dimensions[chr(64 + j)].width = 70 if h in LONG_COLUMNS else 16
        sh.freeze_panes = "A2"
        sh.sheet_properties.tabColor = "6B1D2B"


def write_readme_sheet(wb):
    name = PREFIX + "README"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 130
    for i, line in enumerate(README_LINES, 1):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = WRAP
        c.font = Font(bold=(i == 1), size=14 if i == 1 else 11, color="6B1D2B" if i == 1 else "000000")
    ws.sheet_properties.tabColor = "6B1D2B"


# What the page's engine looks rows up by — a value outside these lists would break the live page.
PLANETS = ["Sun", "Moon", "Mars", "Mercury", "Jupiter", "Venus", "Saturn", "Rahu", "Ketu"]
CLASS_NAMES = ["Kendra", "Trikona", "Panapara", "Apoklima", "Upachaya", "Apachaya", "Maraka", "Dusthana",
               "Badhaka", "Trishadaya"]
STATUSES = ["taught", "curated", "blend", "standard"]
DIGNITY_LABELS = ["Deep Exalted", "Exalted", "Own (Moolatrikona)", "Own House", "Friend's House",
                  "Neutral House", "Enemy's House", "Debilitated", "Deep Debilitated"]
NODE_LABEL = "Node (no rulership)"         # optional row: a dignity text for Rahu/Ketu outside exaltation
DASHA_ROLES = ["Maraka lord", "Maraka occupant", "Badhakadhipati", "Badhaka occupant", "Dusthana lord",
               "Trishadaya lord"]
BADHAKA_MODES = ["Chara", "Sthira", "Dwisabhava"]
NATURES = ["benefic", "malefic", "any"]
# Aspects in the teacher's count (Rahu and Ketu count anti-clockwise: 2nd, 5th, 9th).
ASPECT_COUNTS = {"Sun": [7], "Moon": [7], "Mars": [4, 7, 8], "Mercury": [7], "Jupiter": [5, 7, 9], "Venus": [7],
                 "Saturn": [3, 7, 10], "Rahu": [2, 5, 9], "Ketu": [2, 5, 9]}
# Conditional sentences (S23_Conditions): the engines evaluate exactly these keys (spec §6.6).
CONDITION_KEYS = ("saturn_matures", "saturn_retro_1", "saturn_afflicted_10_young", "saturn_mars_12", "saturn_afflicted_6",
                  "saturn_afflicted_12", "jupiter_md_8", "jupiter_md_11", "venus_dasha_9", "rahu_md_9",
                  "twelfth_hidden_talent", "upachaya_30s", "ketu_12_purpose", "venus_afflicted", "venus_good",
                  "venus_mercury_5", "seventh_lord_12", "mercury_foreign_language", "moon_dual_10", "venus_meets_wife",
                  "jupiter_husband", "first_child_male")
SHEET_ONLY_CONDITION_KEYS = ("twelfth_house",)      # shown by Layer 1 itself (the 12th-house line)
LIFE_AREA_COUNT = 16
LINKS = ["", "PAC"]
META = {"source": "Sessions 23 (2024 and 2026 decks) to 27 slides + recording", "version": 2,
        "statuses": ["taught", "curated", "blend", "standard"]}


class _Sheets:
    """Reads the S23_ data sheets row by row, collecting every problem with its sheet and row number."""

    def __init__(self, wb):
        self.wb, self.problems = wb, []

    def rows(self, short):
        name = PREFIX + short
        width = len(SHEETS[name])
        for i, r in enumerate(self.wb[name].iter_rows(min_row=2, values_only=True), 2):
            r = list(r)[:width] + [None] * max(0, width - len(r))
            if any(v not in (None, "") for v in r):
                yield name, i, [v.strip() if isinstance(v, str) else v for v in r]

    def bad(self, sheet, row, msg):
        self.problems.append(f"{sheet} row {row}: {msg}" if row else f"{sheet}: {msg}")

    def num(self, sheet, row, col, v, lo=1, hi=12):
        try:
            n = float(v)
            if n != int(n):
                raise ValueError
            n = int(n)
        except (TypeError, ValueError):
            self.bad(sheet, row, f"{col} {v!r} is not a whole number")
            return None
        if not lo <= n <= hi:
            self.bad(sheet, row, f"{col} {n} is outside {lo}–{hi}")
            return None
        return n

    def nums(self, sheet, row, col, v):
        out = []
        for part in _s(v).replace(";", ",").split(","):
            if part.strip():
                out.append(self.num(sheet, row, col, part.strip()))
        return out

    def pick(self, sheet, row, col, v, allowed):
        if v not in allowed:
            self.bad(sheet, row, f"{col} {v!r} is not one of {', '.join(x or '(blank)' for x in allowed)}")
        return v

    def opt_num(self, sheet, row, col, v, lo=1, hi=12):
        return None if v in (None, "") else self.num(sheet, row, col, v, lo, hi)

    def opt_pick(self, sheet, row, col, v, allowed):
        return None if v in (None, "") else self.pick(sheet, row, col, v, allowed)

    def text(self, sheet, row, col, v):
        if not _s(v).strip():
            self.bad(sheet, row, f"{col} is empty")
        return _s(v)


def _s(v):
    return "" if v is None else str(v)


def has_session23_sheets(wb):
    return all(name in wb.sheetnames for name in SHEETS)


def _expect(rd, sheet, have, want, what):
    """Each of `want` exactly once in `have` (a list of keys)."""
    for k in want:
        if have.count(k) == 0:
            rd.bad(sheet, None, f"no row for {what(k)}")
        elif have.count(k) > 1:
            rd.bad(sheet, None, f"{have.count(k)} rows for {what(k)} (keep one)")


def read_workbook(path_or_wb):
    """S23_ data sheets -> the rules dict (everything in the JSON except `reference` and `meta`).

    Every cell the page's engine relies on is checked; if anything is wrong, SystemExit lists each
    problem with its sheet and row, and the caller publishes nothing."""
    wb = path_or_wb if hasattr(path_or_wb, "sheetnames") else load_workbook(path_or_wb, data_only=True)
    missing = [n for n in SHEETS if n not in wb.sheetnames]
    if missing:
        raise SystemExit(f"Workbook has no Session 23 sheets {missing}. Run: python make_session23_xlsx.py --install")
    rd = _Sheets(wb)
    out = {}

    out["classes"] = []
    for sh, i, r in rd.rows("Classes"):
        name = rd.pick(sh, i, "Class", r[1], CLASS_NAMES)
        by_lagna = _s(r[3]) in ("", "by Lagna")
        if by_lagna and name != "Badhaka":
            rd.bad(sh, i, f"Houses is empty — only Badhaka is 'by Lagna'")
        out["classes"].append(dict(no=rd.num(sh, i, "No", r[0], 1, len(CLASS_NAMES)), name=name, slide_name=r[2],
                                   houses=None if by_lagna else rd.nums(sh, i, "Houses", r[3]),
                                   other_names=_s(r[4]), nature=_s(r[5]), rule=_s(r[6]), slides=_s(r[7])))
    _expect(rd, PREFIX + "Classes", [c["name"] for c in out["classes"]], CLASS_NAMES, lambda k: f"the class {k}")

    out["class_rules"] = []
    for sh, i, r in rd.rows("ClassRules"):
        out["class_rules"].append({"class": rd.pick(sh, i, "Class", r[0], CLASS_NAMES),
                                   "applies": rd.pick(sh, i, "AppliesTo", r[1], ["any", "benefic", "malefic"]),
                                   "exclude_houses": rd.nums(sh, i, "ExcludeHouses", r[2]),
                                   "text": rd.text(sh, i, "Text", r[3]), "slide": _s(r[4])})
    order = [CLASS_NAMES.index(x["class"]) for x in out["class_rules"] if x["class"] in CLASS_NAMES]
    if order != sorted(order):
        rd.bad(PREFIX + "ClassRules", None, "rows must stay grouped in the order of the ten classes (Kendra … Trishadaya)")

    out["badhaka"] = {}
    for sh, i, r in rd.rows("Badhaka"):
        out["badhaka"][rd.pick(sh, i, "Mode", r[0], BADHAKA_MODES)] = rd.num(sh, i, "House", r[1])
    _expect(rd, PREFIX + "Badhaka", [r[0] for _, _, r in rd.rows("Badhaka")], BADHAKA_MODES, lambda k: f"the {k} mode")

    out["dignity_effect"] = {}
    for sh, i, r in rd.rows("DignityEffect"):
        out["dignity_effect"][rd.pick(sh, i, "Label", r[0], DIGNITY_LABELS + [NODE_LABEL])] = dict(
            band=rd.pick(sh, i, "Band", r[1], ["strong", "medium", "weak"]), text=rd.text(sh, i, "Text", r[2]),
            status=rd.pick(sh, i, "Status", r[3], STATUSES), slide=_s(r[4]))
    _expect(rd, PREFIX + "DignityEffect", [r[0] for _, _, r in rd.rows("DignityEffect")], DIGNITY_LABELS,
            lambda k: f"the dignity {k!r}")

    out["digbala"] = {}
    for sh, i, r in rd.rows("Digbala"):
        out["digbala"][rd.pick(sh, i, "Planet", r[0], PLANETS[:7])] = dict(
            strong=rd.num(sh, i, "StrongHouse", r[1]), lost=rd.num(sh, i, "LostHouse", r[2]),
            status=rd.pick(sh, i, "Status", r[3], STATUSES), source=_s(r[4]))
    _expect(rd, PREFIX + "Digbala", [r[0] for _, _, r in rd.rows("Digbala")], PLANETS[:7], lambda k: k)

    out["dasha_role_text"] = {}
    for sh, i, r in rd.rows("DashaRoleText"):
        out["dasha_role_text"][rd.pick(sh, i, "Role", r[0], DASHA_ROLES)] = dict(
            kind=rd.pick(sh, i, "Kind", r[1], ["caution", "favourable"]), text=rd.text(sh, i, "Text", r[2]), slide=_s(r[3]))
    _expect(rd, PREFIX + "DashaRoleText", [r[0] for _, _, r in rd.rows("DashaRoleText")], DASHA_ROLES, lambda k: f"the role {k!r}")

    out["graha_in_bhava"] = []
    for sh, i, r in rd.rows("GrahaInBhava"):
        points = [p for p in _s(r[2]).split("\n\n") if p.strip()]
        if not points:
            rd.bad(sh, i, "Points is empty")
        out["graha_in_bhava"].append(dict(planet=rd.pick(sh, i, "Planet", r[0], PLANETS), house=rd.num(sh, i, "House", r[1]),
                                          points=points, extra=[p for p in _s(r[3]).split("\n") if p.strip()],
                                          status=rd.pick(sh, i, "Status", r[4], STATUSES), source=_s(r[5])))
    _expect(rd, PREFIX + "GrahaInBhava", [(x["planet"], x["house"]) for x in out["graha_in_bhava"]],
            [(p, h) for p in PLANETS for h in range(1, 13)], lambda k: f"{k[0]} in house {k[1]}")

    out["graha_pair"] = []
    for sh, i, r in rd.rows("GrahaPair"):
        a, b = rd.pick(sh, i, "A", r[0], PLANETS), rd.pick(sh, i, "B", r[1], PLANETS)
        if a == b:
            rd.bad(sh, i, f"A and B are both {a}")
        elif a in PLANETS and b in PLANETS and PLANETS.index(a) > PLANETS.index(b):
            a, b = b, a                       # the engines look pairs up in planet order (Sun … Ketu)
        out["graha_pair"].append(dict(a=a, b=b, conjunction=_s(r[2]), aspect=_s(r[3]),
                                      status=rd.pick(sh, i, "Status", r[4], STATUSES), source=_s(r[5])))
    _expect(rd, PREFIX + "GrahaPair", [frozenset((x["a"], x["b"])) for x in out["graha_pair"]],
            [frozenset((a, b)) for k, a in enumerate(PLANETS) for b in PLANETS[k + 1:]],
            lambda k: " and ".join(sorted(k, key=PLANETS.index)))

    out["bhava_lord_in"] = []
    for sh, i, r in rd.rows("BhavaLordIn"):
        out["bhava_lord_in"].append(dict(lord_of=rd.num(sh, i, "LordOf", r[0]), sits_in=rd.num(sh, i, "SitsIn", r[1]),
                                         text=rd.text(sh, i, "Text", r[2]), status=rd.pick(sh, i, "Status", r[3], STATUSES),
                                         source=_s(r[4]), condition=rd.pick(sh, i, "Condition", _s(r[5]), ["", "strong", "weak"]),
                                         exchange=rd.pick(sh, i, "Exchange", _s(r[6]).lower(), ["", "yes"]) == "yes"))
    out["graha_rashi"] = []
    for sh, i, r in rd.rows("GrahaRashi"):
        out["graha_rashi"].append(dict(planet=rd.pick(sh, i, "Planet", r[0], PLANETS), rashi=rd.num(sh, i, "Rashi", r[1]),
                                       text=rd.text(sh, i, "Text", r[2]), status=rd.pick(sh, i, "Status", r[3], STATUSES),
                                       source=_s(r[4])))
    out["bhava_nature"] = []
    for sh, i, r in rd.rows("BhavaNature"):
        out["bhava_nature"].append(dict(house=rd.opt_num(sh, i, "House", r[0]), nature=rd.pick(sh, i, "Nature", r[1], NATURES),
                                        planet=rd.opt_pick(sh, i, "Planet", r[2], PLANETS), text=rd.text(sh, i, "Text", r[3]),
                                        status=rd.pick(sh, i, "Status", r[4], STATUSES), source=_s(r[5])))
    out["aspect_meaning"] = []
    for sh, i, r in rd.rows("AspectMeaning"):
        planet = rd.pick(sh, i, "Planet", r[0], PLANETS + ["any"])
        aspect = rd.opt_num(sh, i, "Aspect", r[1])
        if aspect is not None and planet in ASPECT_COUNTS and aspect not in ASPECT_COUNTS[planet]:
            rd.bad(sh, i, f"Aspect {aspect} is not one of {planet}'s aspects ({', '.join(map(str, ASPECT_COUNTS[planet]))})")
        out["aspect_meaning"].append(dict(planet=planet, aspect=aspect, from_house=rd.opt_num(sh, i, "FromHouse", r[2]),
                                          text=rd.text(sh, i, "Text", r[3]), status=rd.pick(sh, i, "Status", r[4], STATUSES),
                                          source=_s(r[5])))
    out["life_areas"] = []
    for sh, i, r in rd.rows("LifeAreas"):
        karakas = [k.strip() for k in _s(r[3]).split(",") if k.strip()]
        if not karakas:
            rd.bad(sh, i, "Karaka is empty")
        out["life_areas"].append(dict(no=rd.num(sh, i, "No", r[0], 1, LIFE_AREA_COUNT), area=rd.text(sh, i, "Area", r[1]),
                                      houses=rd.nums(sh, i, "Houses", r[2]),
                                      karakas=[rd.pick(sh, i, "Karaka", k, PLANETS) for k in karakas],
                                      karaka_female=_s(r[4]) and rd.pick(sh, i, "KarakaFemale", _s(r[4]), PLANETS),
                                      link=rd.pick(sh, i, "Link", _s(r[5]), LINKS), source=_s(r[6])))
        if not out["life_areas"][-1]["houses"]:
            rd.bad(sh, i, "Houses is empty")
    _expect(rd, PREFIX + "LifeAreas", [x["no"] for x in out["life_areas"]], list(range(1, LIFE_AREA_COUNT + 1)),
            lambda k: f"area {k}")
    out["conditions"] = []
    for sh, i, r in rd.rows("Conditions"):
        out["conditions"].append(dict(key=rd.pick(sh, i, "Key", r[0], list(CONDITION_KEYS + SHEET_ONLY_CONDITION_KEYS)),
                                      planet=rd.opt_pick(sh, i, "Planet", r[1], PLANETS), house=rd.opt_num(sh, i, "House", r[2]),
                                      text=rd.text(sh, i, "Text", r[3]), status=rd.pick(sh, i, "Status", r[4], STATUSES),
                                      source=_s(r[5])))
    out["remedies"] = []
    for sh, i, r in rd.rows("Remedies"):
        out["remedies"].append(dict(topic=rd.pick(sh, i, "Topic", r[0], PLANETS + ["Tip"]), text=rd.text(sh, i, "Text", r[1]),
                                    source=_s(r[2])))
    for key, sheet, fields in (("bhava_lord_in", "BhavaLordIn", ("lord_of", "sits_in", "condition", "exchange")),
                               ("graha_rashi", "GrahaRashi", ("planet", "rashi")),
                               ("conditions", "Conditions", ("key", "planet", "house"))):
        seen = [tuple(x[f] for f in fields) for x in out[key]]
        for k in sorted({k for k in seen if seen.count(k) > 1}, key=str):
            rd.bad(PREFIX + sheet, None, f"{seen.count(k)} rows for {' / '.join('any' if v is None else str(v) for v in k)} (keep one)")

    if rd.problems:
        raise SystemExit("The Session 23 sheets have {} problem{} — nothing was published:\n  - {}".format(
            len(rd.problems), "" if len(rd.problems) == 1 else "s", "\n  - ".join(rd.problems)))
    return out


# ---------------------------------------------------------------- reference tables
GRAB_JS = r"""
const fs = require('fs'), vm = require('vm');
const src = fs.readFileSync(process.argv[1], 'utf8');
function grab(name) {
  const m = new RegExp('\\b(?:const|let|var)\\s+' + name + '\\s*=\\s*').exec(src);
  if (!m) throw new Error('missing ' + name);
  let i = m.index + m[0].length; const start = i; let depth = 0, q = null;
  for (; i < src.length; i++) {
    const c = src[i], n = src[i + 1];
    if (q) { if (c === '\\') { i++; continue; } if (c === q) q = null; continue; }
    if (c === '/' && n === '/') { while (i < src.length && src[i] !== '\n') i++; continue; }
    if (c === '/' && n === '*') { i = src.indexOf('*/', i + 2) + 1; continue; }
    if (c === '"' || c === "'" || c === '`') { q = c; continue; }
    if ('{[('.includes(c)) depth++;
    else if ('}])'.includes(c)) depth--;
    else if (c === ';' && depth === 0) break;
  }
  // convert Sets inside the context: a Set made in a vm realm is not `instanceof Set` out here
  return vm.runInNewContext('(function(){ var v = ' + src.slice(start, i) + '; return v instanceof Set ? Array.from(v) : v; })()');
}
const out = {};
for (const name of JSON.parse(process.argv[2])) out[name] = grab(name);
console.log(JSON.stringify(out));
"""


def _trim_profile(profile):
    return {planet: {k: prof[k] for k in PROFILE_KEYS if k in prof} for planet, prof in profile.items()}


def extract_reference(index_path=INDEX):
    """The reference tables as the page holds them, read from index.html through node."""
    node = shutil.which("node")
    if not node:
        raise RuntimeError("node is required to read the page's reference tables")
    p = subprocess.run([node, "-e", GRAB_JS, str(index_path), json.dumps(REFERENCE_NAMES)],
                       capture_output=True, text=True, timeout=120)
    if p.returncode != 0:
        raise RuntimeError(f"reference extraction failed: {p.stderr[-600:]}")
    ref = json.loads(p.stdout)
    ref["PLANET_PROFILE"] = _trim_profile(ref["PLANET_PROFILE"])
    return ref


def load_build_data(root=ROOT):
    """Import scripts/build_data.py (the page builder) so its sheet loaders can be reused; None if absent."""
    path = pathlib.Path(root) / "scripts" / "build_data.py"
    if not path.exists():
        return None
    spec = importlib.util.spec_from_file_location("build_data_for_session23", path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


def reference_from_workbook(wb, bd):
    """The reference tables built from the master workbook's own sheets, with build_data.py's loaders —
    the same values scripts/build_data.py writes into index.html."""
    rashi = bd.load_rashi(wb)
    dignity, lord_by_rashi = bd.load_dignity(wb)
    plts = bd.load_planets(wb)
    ref = {
        "BHAVA_INFO": {str(h): v for h, v in bd.load_bhava_info(wb).items()},
        "RASHI": [dict(r, lord=lord_by_rashi.get(r["n"], "")) for r in rashi],
        "RASHI_DIGNITY": {str(n): v for n, v in dignity.items()},
        "DIGNITY_DEG": plts["DIGNITY_DEG"],
        "PLANET_OWN_HOUSES": bd.derive_own_houses(lord_by_rashi),
        "NATURAL_FRIENDS": plts["NATURAL_FRIENDS"],
        "KARAKATWAS": plts["KARAKATWAS"],
        "SPECIAL_ASPECTS": plts["SPECIAL_ASPECTS"],
        "NAKSHATRAS": bd.load_nakshatras(wb)["NAKSHATRAS"],
        "PLANET_PROFILE": _trim_profile(bd.load_planet_profile(wb)),
    }
    ref.update(json.loads(json.dumps(FIXED_REFERENCE)))
    return {name: ref[name] for name in REFERENCE_NAMES}


def reference_for(wb, bd, index_path=INDEX, page_reference=None):
    """The reference tables from the master workbook's sheets when build_data.py can read them there;
    for a workbook without those sheets (or without build_data.py), the tables the page holds."""
    if bd is not None:
        try:
            return reference_from_workbook(wb, bd)
        except KeyError:              # openpyxl: "Worksheet ... does not exist" — not the master workbook
            pass
    return (page_reference or extract_reference)(index_path)


# ---------------------------------------------------------------- build
DATA_START = "/* SESSION23-DATA-START — generated from the S23_ sheets of the master workbook by build_session23.py; do not edit by hand */"
DATA_END = "/* SESSION23-DATA-END */"
ENGINE_START = "/* SESSION23-ENGINE-START"


def page_data_block(rules):
    """The generated `const S23 = {...}` block: the rules without `reference` (the page has those tables itself)."""
    lines = [DATA_START, "const S23 = {"]
    keys = [k for k in rules if k not in ("reference", "meta")]
    for i, k in enumerate(keys):
        body = json.dumps(rules[k], ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
        lines.append(f" {json.dumps(k)}: {body}{',' if i < len(keys) - 1 else ''}")
    lines += ["};", DATA_END]
    return "\n".join(lines) + "\n"


def write_page_block(index_path, rules):
    """Replace (or, the first time, insert before the engine block) the SESSION23-DATA block in index.html."""
    index_path = pathlib.Path(index_path)
    page = index_path.read_text(encoding="utf-8")
    block = page_data_block(rules)
    marker = DATA_START.split(" — ")[0]            # tolerate an older marker comment after the name
    if marker in page:
        a = page.index(marker)
        z = page.index(DATA_END, a) + len(DATA_END) + 1
        new = page[:a] + block + page[z:]
    else:
        if ENGINE_START not in page:
            raise RuntimeError("index.html has no SESSION23-ENGINE block to anchor the data block")
        i = page.index(ENGINE_START)
        new = page[:i] + block + page[i:]
    if new != page:
        index_path.write_text(new, encoding="utf-8")
    return new != page


def find_master(root=ROOT):
    """The master workbook, chosen the way scripts/build_data.py chooses it (v7_1 first, else the newest)."""
    root = pathlib.Path(root)
    preferred = root / "Classification_for_Horoscope_Analysis_v7_1.xlsx"
    if preferred.exists():
        return preferred
    found = sorted(root.glob(MASTER_GLOB))
    return found[-1] if found else None


def build(xlsx_path=None, json_path=JSON_PATH, index_path=INDEX, write_page=False, reference=None, root=ROOT):
    """S23_ sheets of the workbook -> JSON (and, with write_page, the index.html data block).

    `reference` may be passed in (tests); otherwise it comes from the master's own sheets through
    scripts/build_data.py, or — when that script is not there — from index.html through node."""
    xlsx_path = pathlib.Path(xlsx_path) if xlsx_path else find_master(root)
    if not xlsx_path or not xlsx_path.exists():
        raise SystemExit(f"No {MASTER_GLOB} workbook found in {root}")
    json_path = pathlib.Path(json_path)
    wb = load_workbook(xlsx_path, data_only=True)
    rules = read_workbook(wb)
    if reference is None:
        reference = reference_for(wb, load_build_data(root), index_path)
    empty = [k for k in REFERENCE_NAMES if reference.get(k) in (None, "", [], {})]
    if empty:
        raise SystemExit(f"Reference tables came out empty ({', '.join(empty)}) — nothing was published. "
                         "Check the master workbook's sheets and scripts/build_data.py.")
    rules["reference"] = reference
    rules["meta"] = dict(META)
    json_path.write_text(json.dumps(rules, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    if write_page:
        write_page_block(index_path, rules)
    return rules


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    for stream in (sys.stdout, sys.stderr):      # never crash on a character the console's code page lacks
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass
    if "-h" in argv or "--help" in argv:
        print(__doc__)
        return 0
    if "--check" in argv:                       # validate the S23_ sheets only; write nothing
        path = argv[argv.index("--workbook") + 1] if "--workbook" in argv else find_master()
        if not path or not pathlib.Path(path).exists():
            raise SystemExit(f"No {MASTER_GLOB} workbook found in {ROOT}")
        read_workbook(path)
        print("Session 23 sheets: OK")
        return 0
    wb = None
    if "--workbook" in argv:
        wb = argv[argv.index("--workbook") + 1]
    rules = build(wb, write_page=True)
    print(f"Session 23: {len(rules['graha_in_bhava'])} graha-in-bhava rows, {len(rules['graha_pair'])} pairs "
          f"-> {JSON_PATH.name} and the SESSION23-DATA block of {INDEX.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
