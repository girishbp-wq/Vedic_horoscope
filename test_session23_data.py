"""Tests for the Session 23 rules data: workbook -> JSON build (python3 -m unittest -v test_session23_data)."""
import json
import os
import pathlib
import shutil
import subprocess
import unittest

import build_session23 as b23

ROOT = pathlib.Path(__file__).resolve().parent
JSON_PATH = ROOT / "session23_rules.json"
INDEX = ROOT / "index.html"
STATUSES = {"taught", "curated", "blend", "standard"}

CLASS_HOUSES = {
    "Kendra": [1, 4, 7, 10], "Trikona": [1, 5, 9], "Panapara": [2, 5, 8, 11],
    "Apoklima": [3, 6, 9, 12], "Upachaya": [3, 6, 10, 11], "Apachaya": [1, 2, 4, 7, 8],
    "Maraka": [2, 7], "Dusthana": [6, 8, 12], "Badhaka": None, "Trishadaya": [3, 6, 11],
}
PROFILE_KEYS = ["Relationships", "Profession / Status", "Reputation", "Personality Traits",
                "Body Parts Ruled", "General Health Indicator", "Nature", "Combustion Effect",
                "Soul / Mind / Body"]


def rules():
    return json.loads(JSON_PATH.read_text(encoding="utf-8"))


class RuleSheets(unittest.TestCase):
    def test_classes_have_the_confirmed_houses(self):
        got = {c["name"]: c["houses"] for c in rules()["classes"]}
        self.assertEqual(got, CLASS_HOUSES)
        self.assertEqual([c["no"] for c in rules()["classes"]], list(range(1, 11)))

    def test_class_rules_stay_grouped_in_class_order(self):
        # the Excel calculator joins rule texts in sheet order; Python joins them class by class
        order = ["Kendra", "Trikona", "Panapara", "Apoklima", "Upachaya", "Apachaya", "Maraka", "Dusthana", "Badhaka", "Trishadaya"]
        idx = [order.index(r["class"]) for r in rules()["class_rules"]]
        self.assertEqual(idx, sorted(idx))

    def test_trishadaya_malefic_rule_states_its_condition(self):
        # slide 23: amazing results only when the lord of a Trishadaya house is a natural malefic in its own house
        rule = next(x for x in rules()["class_rules"] if x["class"] == "Trishadaya" and x["applies"] == "malefic")
        self.assertIn("own house", rule["text"])
        self.assertIn("lord", rule["text"])

    def test_curated_pair_text_adds_no_named_yogas(self):
        for x in rules()["graha_pair"]:
            if x["status"] == "curated":
                self.assertNotIn("yoga", (x["conjunction"] + " " + x["aspect"]).lower(), (x["a"], x["b"]))

    def test_sun_saturn_and_mars_digbala_are_taught(self):
        self.assertEqual({p: rules()["digbala"][p]["status"] for p in ("Sun", "Saturn", "Mars")},
                         {"Sun": "taught", "Saturn": "taught", "Mars": "taught"})
        self.assertEqual(rules()["digbala"]["Mars"]["source"], "S23-2024 p.29")

    def test_dusthana_malefic_rule_from_session_25(self):
        rows = [x for x in rules()["class_rules"] if x["class"] == "Dusthana" and x["applies"] == "malefic"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["slide"], "S25 p.10")
        import session23_teacher_slides as ts
        self.assertEqual(rows[0]["text"], ts.CLASS_RULE_TEXT["dusthana_malefic"][0])

    def test_apoklima_wording_is_kept_as_printed(self):
        apok = next(c for c in rules()["classes"] if c["name"] == "Apoklima")
        self.assertIn("Except 9th house, it is considered as bad position to the planets", apok["rule"])

    def test_badhaka_modes(self):
        # the page spells the mode "Dwisabhava"; the slides say "Dwiswabhava"
        self.assertEqual(rules()["badhaka"], {"Chara": 11, "Sthira": 9, "Dwisabhava": 7})

    def test_digbala_and_dignity_tables_cover_every_graha_and_label(self):
        r = rules()
        self.assertEqual(sorted(r["digbala"]),
                         sorted(["Sun", "Moon", "Mars", "Mercury", "Jupiter", "Venus", "Saturn"]))
        self.assertEqual((r["digbala"]["Sun"]["strong"], r["digbala"]["Sun"]["lost"]), (10, 4))
        self.assertEqual((r["digbala"]["Saturn"]["strong"]), 7)
        self.assertEqual(sorted(r["dignity_effect"]), sorted(
            ["Deep Exalted", "Exalted", "Own (Moolatrikona)", "Own House", "Friend's House",
             "Neutral House", "Enemy's House", "Debilitated", "Deep Debilitated"]))

    def test_dasha_role_text_has_all_six_roles_in_order(self):
        self.assertEqual(list(rules()["dasha_role_text"]),
                         ["Maraka lord", "Maraka occupant", "Badhakadhipati", "Badhaka occupant",
                          "Dusthana lord", "Trishadaya lord"])
        self.assertIn("before the time of death is promised",
                      rules()["dasha_role_text"]["Maraka lord"]["text"])

    def test_class_rules_cover_the_taught_effects(self):
        cr = rules()["class_rules"]
        has = lambda cls, who, frag: any(x["class"] == cls and x["applies"] == who and frag in x["text"] for x in cr)
        self.assertTrue(has("Kendra", "benefic", "smaller efforts"))
        self.assertTrue(has("Upachaya", "malefic", "good results"))
        self.assertTrue(has("Upachaya", "malefic", "mostly in the 30s"))           # R2: S23 + S26 + S23-2024/S24
        self.assertTrue(has("Apachaya", "malefic", "do not do well"))
        self.assertTrue(has("Panapara", "any", "moderately strong"))
        self.assertTrue(any(x["class"] == "Apoklima" and x["exclude_houses"] == [9] for x in cr))

    def test_sun_is_taught_in_all_twelve_bhavas(self):
        rows = [x for x in rules()["graha_in_bhava"] if x["planet"] == "Sun"]
        self.assertEqual(sorted(x["house"] for x in rows), list(range(1, 13)))
        for x in rows:
            self.assertEqual(x["status"], "taught")
            self.assertTrue(x["points"] and x["extra"], x["house"])

    def test_every_text_row_has_status_and_source(self):
        r = rules()
        for row in r["graha_in_bhava"] + r["graha_pair"]:
            self.assertIn(row["status"], STATUSES, row)
            self.assertTrue(row["source"].strip(), row)
        for sect in ("class_rules", "dasha_role_text"):
            rows = r[sect] if isinstance(r[sect], list) else list(r[sect].values())
            for row in rows:
                self.assertTrue(row.get("slide") or row.get("source"), (sect, row))

    def test_jupiter_mercury_pair_is_taught(self):
        pair = next(x for x in rules()["graha_pair"] if (x["a"], x["b"]) == ("Mercury", "Jupiter"))
        self.assertEqual(pair["status"], "taught")
        self.assertIn("intelligent", pair["conjunction"])


PLANETS = ["Sun", "Moon", "Mars", "Mercury", "Jupiter", "Venus", "Saturn", "Rahu", "Ketu"]


class CuratedContent(unittest.TestCase):
    def test_every_graha_bhava_cell_exists_once(self):
        keys = [(x["planet"], x["house"]) for x in rules()["graha_in_bhava"]]
        self.assertEqual(len(keys), 108)
        self.assertEqual(sorted(set(keys)), sorted((p, h) for p in PLANETS for h in range(1, 13)))

    def test_sun_rows_are_taught_and_equal_the_slide_text(self):
        import session23_sun_slides as slides
        for x in (r for r in rules()["graha_in_bhava"] if r["planet"] == "Sun"):
            self.assertEqual(x["status"], "taught")
            self.assertEqual("\n\n".join(x["points"]), slides.SUN_SLIDE_TEXT[x["house"]][1], x["house"])

    def test_every_graha_bhava_row_is_taught(self):
        import session23_teacher_slides as ts
        for x in rules()["graha_in_bhava"]:
            self.assertEqual(x["status"], "taught", (x["planet"], x["house"]))
            if x["planet"] != "Sun":
                src, text = ts.GRAHA_IN_BHAVA[(x["planet"], x["house"])]
                self.assertEqual(("\n\n".join(x["points"]), x["source"], x["extra"]), (text, src, []), (x["planet"], x["house"]))

    def test_no_two_rows_share_text(self):
        # taught Sun slides legitimately repeat lines ("Every planet has something good and something bad")
        points = [p for x in rules()["graha_in_bhava"] if x["status"] != "taught" for p in x["points"]]
        self.assertEqual(len(points), len(set(points)))
        pairs = [t for x in rules()["graha_pair"] for t in (x["conjunction"], x["aspect"]) if t]
        self.assertEqual(len(pairs), len(set(pairs)))

    def test_curated_rows_cite_their_source_fields(self):
        for x in (r for r in rules()["graha_in_bhava"] if r["status"] == "curated"):
            self.assertIn(f"BHAVA_INFO[{x['house']}:", x["source"], x)
            self.assertTrue(f"PLANET_PROFILE[{x['planet']}:" in x["source"] or f"KARAKATWAS[{x['planet']}:" in x["source"], x)
        for x in (r for r in rules()["graha_pair"] if r["status"] == "curated"):
            self.assertIn("KARAKATWAS[", x["source"])
            self.assertIn(x["a"], x["source"])
            self.assertIn(x["b"], x["source"])

    def test_curated_points_draw_on_the_blended_fields(self):
        import re
        ref = rules()["reference"]
        words = lambda s: set(re.findall(r"[a-z]{5,}", s.lower()))
        for x in (r for r in rules()["graha_in_bhava"] if r["status"] == "curated"):
            b = ref["BHAVA_INFO"][str(x["house"])]
            pool = words(" ".join([b["nm"], b["body"], b["rel"], b["sig"]] + list(ref["KARAKATWAS"][x["planet"]].values())
                                  + [str(v) for v in ref["PLANET_PROFILE"][x["planet"]].values()]))
            pool |= words("self personality health wealth speech family courage communication home mother property "
                          "children creativity intellect enemies debts disease spouse partnerships marriage longevity "
                          "dharma fortune father career status profession gains friends desires losses foreign sleep "
                          "spiritual moksha mind body money savings travel studies education")
            grounded = [pt for pt in x["points"] if words(pt) & pool]
            self.assertGreaterEqual(len(grounded) * 2, len(x["points"]), (x["planet"], x["house"]))

    def test_pair_rows_cover_all_36_and_jupiter_mercury_is_taught(self):
        pairs = rules()["graha_pair"]
        keys = [(x["a"], x["b"]) for x in pairs]
        want = [(a, b) for i, a in enumerate(PLANETS) for b in PLANETS[i + 1:]]
        self.assertEqual(len(keys), 36)
        self.assertEqual(sorted(keys, key=want.index), want)
        for x in pairs:
            self.assertEqual(x["status"], "taught" if (x["a"], x["b"]) == ("Mercury", "Jupiter") else "curated")
            self.assertTrue(x["conjunction"].strip())
            if x["status"] == "curated":
                self.assertTrue(x["aspect"].strip(), (x["a"], x["b"]))

    def test_tone_has_no_certainty_words(self):
        import re
        banned = re.compile(r"will die|certainly|definitely|surely|guaranteed|inevitabl|cannot fail", re.I)
        r = rules()
        texts = [p for x in r["graha_in_bhava"] if x["status"] == "curated" for p in x["points"]]
        texts += [t for x in r["graha_pair"] if x["status"] == "curated" for t in (x["conjunction"], x["aspect"])]
        for t in texts:
            self.assertIsNone(banned.search(t), t)


class TeacherRows(unittest.TestCase):
    """Sessions 23 (2024 deck) to 27: the rows come from session23_teacher_slides.py."""

    @classmethod
    def _exported(cls):
        import tempfile
        import make_session23_xlsx as mx
        cls._tmp = tempfile.TemporaryDirectory()
        path = pathlib.Path(cls._tmp.name, "export.xlsx")
        mx.build_workbook(path)
        return path

    def test_new_keys_equal_the_teacher_module(self):
        import session23_teacher_slides as ts
        r = rules()
        self.assertEqual(r["bhava_nature"], [dict(house=h, nature=n, planet=p, text=t, status="taught", source=s)
                                             for h, n, p, t, s in ts.BHAVA_NATURE])
        self.assertEqual(r["aspect_meaning"], [dict(planet=p, aspect=a, from_house=f, text=t, status="taught", source=s)
                                               for p, a, f, t, s in ts.ASPECT_MEANING])
        self.assertEqual(r["life_areas"], [dict(no=n, area=a, houses=list(h), karakas=list(k), karaka_female=kf, link=l, source=s)
                                           for n, a, h, k, kf, l, s in ts.LIFE_AREAS])
        self.assertEqual(r["conditions"], [dict(key=k, planet=p, house=h, text=t, status="taught", source=s)
                                           for k, p, h, t, s in ts.CONDITIONS])
        self.assertEqual(r["remedies"], [dict(topic=t, text=x, source=s) for t, x, s in ts.REMEDIES])

    def test_lord_in_and_graha_rashi_rows_follow_the_shipped_ones(self):
        import session23_teacher_slides as ts
        r = rules()
        self.assertEqual(r["bhava_lord_in"][0]["lord_of"], 2)                  # the Session 23 row stays first
        self.assertEqual(r["bhava_lord_in"][1:], [dict(lord_of=a, sits_in=b, text=t, status="taught", source=s, condition=c, exchange=x)
                                                  for a, b, c, x, t, s in ts.LORD_IN])
        self.assertEqual(r["graha_rashi"][1:], [dict(planet=p, rashi=n, text=t, status="taught", source=s)
                                                for p, n, t, s in ts.GRAHA_RASHI])
        self.assertEqual((r["bhava_lord_in"][0]["condition"], r["bhava_lord_in"][0]["exchange"]), ("", False))

    def test_json_is_previous_plus_teacher_slides(self):
        import apply_teacher_slides
        prev = json.loads((ROOT / "session23_rules.previous.json").read_text(encoding="utf-8"))
        # the reference tables come from the master workbook's own sheets, not from the slides
        drop = lambda r: {k: v for k, v in r.items() if k != "reference"}
        self.assertEqual(drop(apply_teacher_slides.apply(prev)), drop(rules()))
        # key order too: the workbook build and this script must write the same file, so publish sees no change
        built = b23.read_workbook(self.__class__._exported())
        for k, rows in built.items():
            if isinstance(rows, list):
                self.assertEqual([list(r) for r in rows], [list(r) for r in apply_teacher_slides.apply(prev)[k]], k)
        self.assertEqual(prev["meta"]["version"], 1)
        self.assertEqual(rules()["meta"], b23.META)


class TaughtLayerExamples(unittest.TestCase):
    def test_second_lord_in_seventh_is_taught_from_the_recording(self):
        rows = [x for x in rules()["bhava_lord_in"] if (x["lord_of"], x["sits_in"]) == (2, 7)]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["status"], "taught")
        for frag in ("family business", "source of income", "income from your spouse"):
            self.assertIn(frag, rows[0]["text"])
        self.assertTrue(rows[0]["source"].strip())

    def test_sun_in_mesha_is_taught_from_the_recording(self):
        rows = [x for x in rules()["graha_rashi"] if (x["planet"], x["rashi"]) == ("Sun", 1)]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["status"], "taught")
        for frag in ("Agni", "eastern", "Kshatriya", "courage", "hilly", "exalted", "full potential"):
            self.assertIn(frag, rows[0]["text"])


class WorkbookRoundTrip(unittest.TestCase):
    """The S23_ sheets are the source of truth; these tests check that writing them and reading them back is lossless."""

    @classmethod
    def setUpClass(cls):
        import tempfile
        import make_session23_xlsx as mx
        cls.tmp = tempfile.TemporaryDirectory()
        cls.xlsx = pathlib.Path(cls.tmp.name, "export.xlsx")
        mx.build_workbook(cls.xlsx)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_json_matches_an_exported_workbook(self):
        from_xlsx = b23.read_workbook(self.xlsx)
        from_json = {k: v for k, v in rules().items() if k not in ("reference", "meta")}
        self.assertEqual(from_xlsx, from_json)

    def test_exported_workbook_has_the_prefixed_sheets(self):
        import openpyxl
        names = openpyxl.load_workbook(self.xlsx).sheetnames
        for n in ("README", "Classes", "ClassRules", "Badhaka", "DignityEffect", "Digbala", "DashaRoleText",
                  "GrahaInBhava", "GrahaPair", "BhavaLordIn", "GrahaRashi", "BhavaNature", "AspectMeaning",
                  "LifeAreas", "Conditions", "Remedies", "Chart", "Classes_Calc", "Roles_Calc",
                  "Dasha_Calc", "Predict_Calc", "Ref_Calc"):
            self.assertIn("S23_" + n, names)

    def test_build_reproduces_the_json_from_a_workbook(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            j = pathlib.Path(d, "r.json")
            b23.build(self.xlsx, j, INDEX, reference=rules()["reference"])
            self.assertEqual(json.loads(j.read_text(encoding="utf-8")), rules())

    def test_reference_is_assembled_from_the_master_sheets_with_the_page_builders_loaders(self):
        ref = rules()["reference"]

        class BuildData:       # stands in for scripts/build_data.py; returns what its loaders return for this workbook
            load_rashi = staticmethod(lambda wb: [{k: v for k, v in r.items() if k != "lord"} for r in ref["RASHI"]])
            load_dignity = staticmethod(lambda wb: ({int(n): v for n, v in ref["RASHI_DIGNITY"].items()},
                                                    {r["n"]: r["lord"] for r in ref["RASHI"]}))
            load_planets = staticmethod(lambda wb: {k: ref[k] for k in ("DIGNITY_DEG", "NATURAL_FRIENDS", "KARAKATWAS", "SPECIAL_ASPECTS")})
            derive_own_houses = staticmethod(lambda lords: ref["PLANET_OWN_HOUSES"])
            load_bhava_info = staticmethod(lambda wb: {int(h): v for h, v in ref["BHAVA_INFO"].items()})
            load_nakshatras = staticmethod(lambda wb: {"NAKSHATRAS": ref["NAKSHATRAS"]})
            load_planet_profile = staticmethod(lambda wb: ref["PLANET_PROFILE"])

        self.assertEqual(b23.reference_from_workbook(None, BuildData), ref)

    def test_master_workbook_is_found_the_way_build_data_finds_it(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            for name in ("Classification_for_Horoscope_Analysis_v7_1.before-session23-2026-10-08.xlsx",
                         "Classification_for_Horoscope_Analysis_v8.xlsx", "Classification_for_Horoscope_Analysis_v7_1.xlsx"):
                pathlib.Path(d, name).write_bytes(b"x")
            self.assertEqual(b23.find_master(d).name, "Classification_for_Horoscope_Analysis_v7_1.xlsx")
            pathlib.Path(d, "Classification_for_Horoscope_Analysis_v7_1.xlsx").unlink()
            self.assertEqual(b23.find_master(d).name, "Classification_for_Horoscope_Analysis_v8.xlsx")

    def test_publish_bat_builds_session23_after_the_page_data_and_stages_the_json(self):
        bat = (ROOT / "publish.bat").read_text(encoding="utf-8")
        self.assertLess(bat.index("scripts\\build_data.py"), bat.index("build_session23.py"))
        self.assertIn("git add index.html session23_rules.json", bat)


class PublishScript(unittest.TestCase):
    """publish.bat runs on Windows cmd; these pin the behaviour the audit asked for."""

    @classmethod
    def setUpClass(cls):
        cls.bat = (ROOT / "publish.bat").read_text(encoding="utf-8")

    def test_only_publishes_from_main_after_pulling_it(self):
        self.assertIn("git rev-parse --abbrev-ref HEAD", self.bat)
        self.assertIn('if /i not "%BRANCH%"=="main"', self.bat)
        self.assertLess(self.bat.index("git pull --ff-only"), self.bat.index("%PY% scripts\\build_data.py"))

    def test_commits_only_its_two_files(self):
        self.assertIn("git diff --cached --quiet -- index.html session23_rules.json", self.bat)
        commit = next(l for l in self.bat.splitlines() if l.strip().startswith("git commit"))
        self.assertTrue(commit.rstrip().endswith("-- index.html session23_rules.json"), commit)

    def test_pushes_an_earlier_commit_whose_push_failed(self):
        self.assertIn("git rev-list --count @{u}..HEAD", self.bat)
        nochange = self.bat.index(":nochange")
        self.assertLess(nochange, self.bat.index(":push"))

    def test_no_wmic_and_plain_ascii(self):
        self.assertNotIn("wmic", self.bat.lower())
        self.assertTrue(self.bat.isascii())
        self.assertNotIn("pause\nexit /b 1\n)", self.bat)       # every failure goes through :fail

    def test_runs_from_a_copy_so_a_pull_cannot_rewrite_it_mid_run(self):
        head = self.bat[:self.bat.index("setlocal")]
        self.assertIn('if /i not "%~1"=="--from-copy" (', head)
        self.assertIn('copy /y "%~f0" "%TEMP%\\jyotisha_publish.bat" >nul', head)
        self.assertIn('call "%TEMP%\\jyotisha_publish.bat" --from-copy "%~dp0."', head)
        self.assertIn('cd /d "%~2"', self.bat)
        self.assertLess(self.bat.index('cd /d "%~2"'), self.bat.index("git pull --ff-only"))

    def test_checks_the_session23_sheets_before_rebuilding_the_page(self):
        self.assertLess(self.bat.index("%PY% build_session23.py --check"), self.bat.index("%PY% scripts\\build_data.py"))

    def test_finds_a_python_that_runs_not_the_store_shortcut(self):
        # Windows has a "python" Store shortcut even with no Python installed: `where python` finds it
        self.assertNotIn("where python", self.bat)
        self.assertIn('python -c "import sys" >nul 2>nul', self.bat)
        self.assertIn('py -3 -c "import sys" >nul 2>nul', self.bat)
        self.assertIn("Add python.exe to PATH", self.bat)
        runs = [l.strip() for l in self.bat.splitlines() if ".py" in l or "-m pip" in l or "import openpyxl" in l]
        runs = [l for l in runs if not l.upper().startswith(("REM", "ECHO"))]
        self.assertTrue(runs)
        for line in runs:
            self.assertTrue(line.startswith("%PY% "), line)
        self.assertLess(self.bat.index('set "PY=py -3"'), self.bat.index("%PY% "))

    def test_checked_out_with_windows_line_endings(self):
        attrs = (ROOT / ".gitattributes").read_text(encoding="utf-8")
        self.assertIn("*.bat   text eol=crlf", attrs)


class SheetValidation(unittest.TestCase):
    """A typo in an S23_ sheet stops the build with a message naming the sheet and row, and nothing is written."""

    def build_after(self, edit):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            wb = openpyxl.Workbook()
            b23.write_data_sheets(wb, rules())
            edit(wb)
            x, j, page = pathlib.Path(d, "w.xlsx"), pathlib.Path(d, "r.json"), pathlib.Path(d, "index.html")
            wb.save(x)
            page.write_text(INDEX.read_text(encoding="utf-8"), encoding="utf-8")
            with self.assertRaises(SystemExit) as cm:
                b23.build(x, j, page, write_page=True, reference=rules()["reference"])
            self.assertFalse(j.exists())
            self.assertEqual(page.read_text(encoding="utf-8"), INDEX.read_text(encoding="utf-8"))
            return str(cm.exception.code)

    @staticmethod
    def find(wb, sheet, col, value):
        ws = wb["S23_" + sheet]
        return next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=col).value == value)

    def test_misspelt_planet_names_sheet_and_row(self):
        def edit(wb):
            r = self.find(wb, "GrahaInBhava", 1, "Saturn")
            wb["S23_GrahaInBhava"].cell(row=r, column=1, value="Saturnn")
            self.row = r
        msg = self.build_after(edit)
        self.assertIn(f"S23_GrahaInBhava row {self.row}: Planet 'Saturnn'", msg)
        self.assertIn("nothing was published", msg)
        self.assertIn("S23_GrahaInBhava: no row for Saturn in house 1", msg)

    def test_house_that_is_not_a_number(self):
        msg = self.build_after(lambda wb: wb["S23_GrahaInBhava"].cell(row=5, column=2, value="x"))
        self.assertIn("S23_GrahaInBhava row 5: House 'x' is not a whole number", msg)

    def test_house_out_of_range_and_bad_status(self):
        def edit(wb):
            wb["S23_Digbala"].cell(row=2, column=2, value=13)
            wb["S23_GrahaPair"].cell(row=3, column=5, value="taugth")
        msg = self.build_after(edit)
        self.assertIn("S23_Digbala row 2: StrongHouse 13 is outside 1–12", msg)
        self.assertIn("S23_GrahaPair row 3: Status 'taugth' is not one of taught, curated, blend, standard", msg)

    def test_missing_pair_and_renamed_class(self):
        def edit(wb):
            wb["S23_GrahaPair"].delete_rows(2)
            wb["S23_Classes"].cell(row=2, column=2, value="Kendraa")
        msg = self.build_after(edit)
        self.assertRegex(msg, r"S23_GrahaPair: no row for \w+ and \w+")
        self.assertIn("S23_Classes row 2: Class 'Kendraa'", msg)

    def test_pair_typed_in_reverse_order_is_stored_in_planet_order(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            wb = openpyxl.Workbook()
            b23.write_data_sheets(wb, rules())
            ws = wb["S23_GrahaPair"]
            r = self.find(wb, "GrahaPair", 1, "Sun")
            a, b = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
            ws.cell(row=r, column=1, value=b)
            ws.cell(row=r, column=2, value=a)
            x = pathlib.Path(d, "w.xlsx")
            wb.save(x)
            self.assertEqual(b23.read_workbook(x)["graha_pair"], rules()["graha_pair"])

    def test_a_node_dignity_text_may_be_added(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            wb = openpyxl.Workbook()
            b23.write_data_sheets(wb, rules())
            ws = wb["S23_DignityEffect"]
            ws.append(["Node (no rulership)", "medium", "gives the results of its sign's lord.", "standard", ""])
            x = pathlib.Path(d, "w.xlsx")
            wb.save(x)
            self.assertIn("Node (no rulership)", b23.read_workbook(x)["dignity_effect"])

    def test_new_sheet_bad_values_are_named(self):
        def edit(wb):
            wb["S23_BhavaNature"].cell(row=2, column=2, value="benefik")
            wb["S23_AspectMeaning"].cell(row=self.find(wb, "AspectMeaning", 1, "Saturn"), column=2, value=5)
            wb["S23_LifeAreas"].cell(row=3, column=4, value="Jupitor")
            wb["S23_Conditions"].cell(row=2, column=1, value="saturn_maturs")
            wb["S23_Remedies"].cell(row=2, column=1, value="Moom")
            wb["S23_BhavaLordIn"].cell(row=2, column=6, value="strng")
            wb["S23_BhavaLordIn"].cell(row=3, column=7, value="maybe")
            self.saturn_row = self.find(wb, "AspectMeaning", 1, "Saturn")
        msg = self.build_after(edit)
        self.assertIn("S23_BhavaNature row 2: Nature 'benefik' is not one of benefic, malefic, any", msg)
        self.assertIn(f"S23_AspectMeaning row {self.saturn_row}: Aspect 5 is not one of Saturn's aspects (3, 7, 10)", msg)
        self.assertIn("S23_LifeAreas row 3: Karaka 'Jupitor'", msg)
        self.assertIn("S23_Conditions row 2: Key 'saturn_maturs'", msg)
        self.assertIn("S23_Remedies row 2: Topic 'Moom'", msg)
        self.assertIn("S23_BhavaLordIn row 2: Condition 'strng' is not one of (blank), strong, weak", msg)
        self.assertIn("S23_BhavaLordIn row 3: Exchange 'maybe' is not one of (blank), yes", msg)

    def test_life_areas_must_be_1_to_16(self):
        msg = self.build_after(lambda wb: wb["S23_LifeAreas"].delete_rows(5))
        self.assertIn("S23_LifeAreas: no row for area 4", msg)

    def test_life_area_cells_the_excel_sheet_cannot_read(self):
        def edit(wb):
            ws = wb["S23_LifeAreas"]
            ws.cell(row=2, column=3, value="1, 2, 3")                  # three houses
            ws.cell(row=3, column=3, value="2; 11")                    # a semicolon
            ws.cell(row=4, column=4, value="Moon, Venus, Sun")         # three karakas
            ws.cell(row=17, column=3, value="5")                       # Love marriage (PAC) with one house
        msg = self.build_after(edit)
        self.assertIn("S23_LifeAreas row 2: Houses '1, 2, 3' — at most two houses", msg)
        self.assertIn("S23_LifeAreas row 3: Houses '2; 11' — separate the houses with a comma", msg)
        self.assertIn("S23_LifeAreas row 4: Karaka 'Moon, Venus, Sun' — at most two karakas", msg)
        self.assertIn("S23_LifeAreas row 17: Link PAC needs two houses and two karakas", msg)

    def test_duplicate_condition_row(self):
        def edit(wb):
            ws = wb["S23_Conditions"]
            ws.append([c.value for c in ws[2]])
        msg = self.build_after(edit)
        self.assertRegex(msg, r"S23_Conditions: 2 rows for \w+")

    def test_duplicate_aspect_meaning_row(self):
        def edit(wb):
            ws = wb["S23_AspectMeaning"]
            ws.append([c.value for c in ws[2]])
        self.assertRegex(self.build_after(edit), r"S23_AspectMeaning: 2 rows for \w+")

    def test_parivartana_typed_both_ways(self):
        def edit(wb):
            ws = wb["S23_BhavaLordIn"]
            r = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=7).value == "yes")
            row = [c.value for c in ws[r]]
            row[0], row[1] = row[1], row[0]
            ws.append(row)
        self.assertIn("S23_BhavaLordIn: 2 Parivartana rows for houses 4 and 11 (keep one)", self.build_after(edit))

    def test_old_bhava_lord_in_sheet_reads_with_blank_new_columns(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            wb = openpyxl.Workbook()
            b23.write_data_sheets(wb, rules())
            ws = wb["S23_BhavaLordIn"]
            ws.delete_cols(6, 2)                                   # the sheet as the previous release wrote it
            for r in range(ws.max_row, 2, -1):
                ws.delete_rows(r)
            x = pathlib.Path(d, "w.xlsx")
            wb.save(x)
            self.assertEqual(b23.read_workbook(x)["bhava_lord_in"], rules()["bhava_lord_in"][:1])

    def test_check_only_reads_and_writes_nothing(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            wb = openpyxl.Workbook()
            b23.write_data_sheets(wb, rules())
            wb["S23_GrahaInBhava"].cell(row=5, column=2, value="x")
            x = pathlib.Path(d, "Classification_for_Horoscope_Analysis_v7_1.xlsx")
            wb.save(x)
            with self.assertRaises(SystemExit) as cm:
                b23.main(["--check", "--workbook", str(x)])
            self.assertIn("S23_GrahaInBhava row 5", str(cm.exception.code))
            self.assertEqual(sorted(p.name for p in pathlib.Path(d).iterdir()), [x.name])
            wb["S23_GrahaInBhava"].cell(row=5, column=2, value=4)           # the Sun in the 4th again
            wb.save(x)
            before = (JSON_PATH.read_bytes(), INDEX.read_bytes())
            self.assertEqual(b23.main(["--check", "--workbook", str(x)]), 0)
            self.assertEqual((JSON_PATH.read_bytes(), INDEX.read_bytes()), before)
            self.assertEqual(sorted(p.name for p in pathlib.Path(d).iterdir()), [x.name])

    def test_stray_spaces_are_tolerated(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            wb = openpyxl.Workbook()
            b23.write_data_sheets(wb, rules())
            wb["S23_GrahaInBhava"].cell(row=2, column=1, value=wb["S23_GrahaInBhava"].cell(row=2, column=1).value + " ")
            x = pathlib.Path(d, "w.xlsx")
            wb.save(x)
            self.assertEqual(b23.read_workbook(x), {k: v for k, v in rules().items() if k not in ("reference", "meta")})

    def test_empty_reference_is_refused(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            wb = openpyxl.Workbook()
            b23.write_data_sheets(wb, rules())
            x, j = pathlib.Path(d, "w.xlsx"), pathlib.Path(d, "r.json")
            wb.save(x)
            ref = dict(rules()["reference"], NAKSHATRAS=[])
            with self.assertRaises(SystemExit) as cm:
                b23.build(x, j, INDEX, reference=ref)
            self.assertIn("NAKSHATRAS", str(cm.exception.code))
            self.assertFalse(j.exists())

    def test_workbook_without_master_sheets_takes_the_reference_from_the_page(self):
        class BuildData:           # a build_data.py whose loaders find no master sheets in this workbook
            def __getattr__(self, name):
                def loader(wb, *a):
                    raise KeyError("Worksheet Reference Data does not exist.")
                return loader
        page_ref = rules()["reference"]
        got = b23.reference_for(object(), BuildData(), INDEX, page_reference=lambda path: page_ref)
        self.assertEqual(got, page_ref)


NODE = shutil.which("node")
NPM = shutil.which("npm")


def _pw():
    if not (NODE and NPM):
        return None
    root = subprocess.run([NPM, "root", "-g"], capture_output=True, text=True).stdout.strip()
    p = pathlib.Path(root) / "playwright"
    return str(p) if p.exists() and os.path.isdir(os.environ.get("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")) else None


PAGE_SCRIPT = r"""
const { chromium } = require(process.env.PW_MODULE);
(async () => {
  const b = await chromium.launch({headless: true});
  const ctx = await b.newContext();
  await ctx.route('**/*', r => r.request().url().startsWith('file:') ? r.continue() : r.abort());
  const p = await ctx.newPage();
  await p.goto('file://' + process.argv[1]);
  const out = await p.evaluate(() => ({
    BHAVA_INFO, RASHI, RASHI_DIGNITY, DIGNITY_DEG, PLANET_OWN_HOUSES, NATURAL_FRIENDS, KARAKATWAS,
    SPECIAL_ASPECTS, VORDER, VYEARS, COMBUST_ORB, COMBUST_PLANETS: [...COMBUST_PLANETS], NAKSHATRAS,
    PLANET_PROFILE
  }));
  console.log(JSON.stringify(out));
  await b.close();
})();
"""


@unittest.skipUnless(_pw(), "playwright + chromium needed to read the live page")
class ReferenceMatchesThePage(unittest.TestCase):
    def test_reference_equals_the_values_the_page_itself_holds(self):
        env = dict(os.environ, PW_MODULE=_pw(), PLAYWRIGHT_BROWSERS_PATH="/opt/pw-browsers")
        p = subprocess.run([NODE, "-e", PAGE_SCRIPT, str(INDEX)], capture_output=True, text=True,
                           timeout=120, env=env)
        self.assertEqual(p.returncode, 0, p.stderr[-800:])
        live = json.loads(p.stdout.strip().splitlines()[-1])
        ref = rules()["reference"]
        for key in ("BHAVA_INFO", "RASHI", "RASHI_DIGNITY", "DIGNITY_DEG", "PLANET_OWN_HOUSES",
                    "NATURAL_FRIENDS", "KARAKATWAS", "SPECIAL_ASPECTS", "VORDER", "VYEARS",
                    "COMBUST_ORB", "COMBUST_PLANETS", "NAKSHATRAS"):
            self.assertEqual(ref[key], live[key], key)
        for planet, prof in live["PLANET_PROFILE"].items():
            want = {k: prof[k] for k in PROFILE_KEYS if k in prof}
            self.assertEqual(ref["PLANET_PROFILE"][planet], want, planet)


if __name__ == "__main__":
    unittest.main()
