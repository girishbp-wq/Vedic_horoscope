"""Cross-implementation tests for Session 23: the page's JavaScript vs bhava_rules.py.

python3 -m unittest -v test_session23_cross_impl
The JS runs in headless Chromium against the real index.html (all network blocked); results are compared to
the Python engine on random charts by canonical-JSON fingerprint, and a mismatch is reported with the first
differing path."""
import json
import os
import pathlib
import random
import re
import shutil
import subprocess
import unittest

import bhava_rules as br

ROOT = pathlib.Path(__file__).resolve().parent
INDEX = ROOT / "index.html"
JSON_PATH = ROOT / "session23_rules.json"
RULES = br.load_rules()
NODE, NPM = shutil.which("node"), shutil.which("npm")


def _pw():
    if not (NODE and NPM):
        return None
    root = subprocess.run([NPM, "root", "-g"], capture_output=True, text=True).stdout.strip()
    p = pathlib.Path(root) / "playwright"
    return str(p) if p.exists() and os.path.isdir(os.environ.get("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")) else None


PW = _pw()

# --------------------------------------------------------------------------------------- canonical form
def _js_numbers(o):
    """JSON.stringify writes 15.0 as 15: integral floats become ints so both sides serialise alike."""
    if isinstance(o, float) and o.is_integer():
        return int(o)
    if isinstance(o, dict):
        return {k: _js_numbers(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_js_numbers(v) for v in o]
    return o


def canon(o):
    return json.dumps(_js_numbers(o), sort_keys=True, ensure_ascii=False, separators=(",", ":"))


def fingerprint(o):
    s = canon(o)
    h = 0x811C9DC5
    b = s.encode("utf-16-le")
    for i in range(0, len(b), 2):
        h = ((h ^ (b[i] | (b[i + 1] << 8))) * 0x01000193) & 0xFFFFFFFF
    return [len(s), h]


NODE_RUNNER = r"""
const { chromium } = require(process.env.PW_MODULE);
(async () => {
  const job = JSON.parse(require('fs').readFileSync(process.argv[2], 'utf8'));
  const b = await chromium.launch({headless: true});
  const ctx = await b.newContext({timezoneId: 'UTC'});
  await ctx.route('**/*', r => r.request().url().startsWith('file:') ? r.continue() : r.abort());
  const p = await ctx.newPage();
  const errors = [];
  p.on('pageerror', e => errors.push(String(e)));
  await p.goto('file://' + job.index);
  const out = await p.evaluate(new Function('job', job.body), job.arg);
  console.log(JSON.stringify({out, errors}));
  await b.close();
})();
"""


def run_page(body, arg, tmp):
    """Run JS `body` (a function body receiving `job`) on the loaded page; returns the JSON-able result."""
    job = tmp / "job.json"
    runner = tmp / "runner.js"
    job.write_text(json.dumps({"index": str(INDEX), "body": body, "arg": arg}), encoding="utf-8")
    runner.write_text(NODE_RUNNER, encoding="utf-8")
    env = dict(os.environ, PW_MODULE=PW, PLAYWRIGHT_BROWSERS_PATH="/opt/pw-browsers")
    p = subprocess.run([NODE, str(runner), str(job)], capture_output=True, text=True, timeout=600, env=env)
    if p.returncode != 0:
        raise AssertionError(p.stderr[-1500:])
    res = json.loads(p.stdout.strip().splitlines()[-1])
    if res["errors"]:
        raise AssertionError("page errors: " + "; ".join(res["errors"]))
    return res["out"]


JS_CANON = r"""
function canon(v){
  if (Array.isArray(v)) return '[' + v.map(canon).join(',') + ']';
  if (v === undefined) return '"__undefined__"';
  if (v && typeof v === 'object') return '{' + Object.keys(v).sort().map(k => JSON.stringify(k) + ':' + canon(v[k])).join(',') + '}';
  return JSON.stringify(v);
}
function fp(v){ const s = canon(v); let h = 0x811c9dc5; for (let i = 0; i < s.length; i++) { h ^= s.charCodeAt(i); h = Math.imul(h, 0x01000193) >>> 0; } return [s.length, h]; }
const PL = ["Sun","Moon","Mars","Mercury","Jupiter","Venus","Saturn","Rahu","Ketu"];
function layers(c){
  const lg = c.lagna + 1, pbr = {}; for (const p of PL) pbr[p] = c.signs[p] + 1;
  return {
    tables: s23ClassificationTables(lg, pbr), matrix: s23HouseClassMatrix(lg), badhaka: s23Badhaka(lg, pbr),
    roles: s23Roles(lg, pbr), readings: s23ClassReadings(lg, pbr),
    graha_bhava: PL.map(p => s23GrahaBhava(lg, pbr, p, c.degs[p], c.waxing)),
    bhava_bhava: s23BhavaBhava(ctxOf(c, lg, pbr)), graha_rashi: s23GrahaRashi(lg, pbr, c.degs),
    graha_graha: s23GrahaGraha(lg, pbr, c.degs),
    five_step: [1,2,3,4,5,6,7,8,9,10,11,12].map(h => s23FiveStep(lg, pbr, h)),
    ...extra(c, lg, pbr),
  };
}
function ctxOf(c, lg, pbr){
  return s23Ctx(lg, pbr, c.degs, c.waxing, {retro: c.retro, gender: c.gender, age: c.age, maha: c.maha, bhukti: c.bhukti});
}
function extra(c, lg, pbr){
  const ctx = ctxOf(c, lg, pbr), placed = PL.filter(p => ctx.signs[p] !== undefined);
  return {
    context: ctx,
    helpers: Object.fromEntries(placed.map(p => [p, {combust: s23Combust(ctx, p), afflicted: s23Afflicted(ctx, p),
      strength: s23Strength(ctx, p), lords_houses: s23LordsHouses(ctx.lagna, p),
      aspects_on: s23AspectsOn(ctx, ctx.signs[p])}])),
    sign_lines: [0,1,2,3,4,5,6,7,8,9,10,11].map(s23SignLine),
    male_signs: [0,1,2,3,4,5,6,7,8,9,10,11].map(s23IsMaleSign),
    conditions: s23ChartConditions(ctx), condition_keys: S23_CONDITION_KEYS,
    life_areas: s23LifeAreas(ctx),
  };
}
"""


def python_layers(c):
    lg, sg, degs, w = c["lagna"], c["signs"], c["degs"], c["waxing"]
    return {
        "tables": br.classification_tables(RULES, lg, sg), "matrix": br.house_class_matrix(RULES, lg),
        "badhaka": br.badhaka(RULES, lg, sg), "roles": br.planet_roles(RULES, lg, sg),
        "readings": br.class_readings(RULES, lg, sg),
        "graha_bhava": [br.graha_bhava(RULES, lg, sg, p, degs[p], w) for p in br.PLANET_ORDER],
        "bhava_bhava": br.bhava_bhava(RULES, ctx_of(c)), "graha_rashi": br.graha_rashi(RULES, lg, sg, degs),
        "graha_graha": br.graha_graha(RULES, lg, sg, degs),
        "five_step": [br.five_step(RULES, lg, sg, h) for h in range(1, 13)],
        **python_extra(c),
    }


def ctx_of(c):
    return br.context(c["lagna"], c["signs"], c["degs"], c["waxing"], c.get("retro"), c.get("gender"), c.get("age"),
                      c.get("maha"), c.get("bhukti"))


def python_extra(c):
    ctx = ctx_of(c)
    placed = [p for p in br.PLANET_ORDER if p in ctx["signs"]]
    return {
        "context": ctx,
        "helpers": {p: {"combust": br.combust(RULES, ctx, p), "afflicted": br.afflicted(RULES, ctx, p),
                        "strength": br.strength(RULES, ctx, p), "lords_houses": br.lords_houses(RULES, ctx["lagna"], p),
                        "aspects_on": br.aspects_on(RULES, ctx, ctx["signs"][p])} for p in placed},
        "sign_lines": [br.sign_line(RULES, s) for s in range(12)],
        "male_signs": [br.is_male_sign(RULES, s) for s in range(12)],
        "conditions": br.chart_conditions(RULES, ctx), "condition_keys": list(br.CONDITION_KEYS),
        "life_areas": br.life_areas(RULES, ctx),
    }


def random_charts(n, seed):
    rnd = random.Random(seed)
    charts = []
    for _ in range(n):
        pool = rnd.sample(range(12), rnd.choice([3, 5, 12]))
        signs = {p: rnd.choice(pool) for p in br.PLANET_ORDER}
        charts.append({"lagna": rnd.randrange(12), "signs": signs, "waxing": rnd.random() < 0.5,
                       "degs": {p: round(rnd.uniform(0, 29.99), 2) for p in br.PLANET_ORDER},
                       "retro": {p: rnd.random() < 0.3 for p in br.PLANET_ORDER},
                       "gender": rnd.choice([None, "Male", "Female", "Other"]),
                       "age": rnd.choice([None, round(rnd.uniform(5, 80), 2)]),
                       "maha": rnd.choice([None] + br.PLANET_ORDER), "bhukti": rnd.choice([None] + br.PLANET_ORDER)})
    return charts


def first_difference(a, b, path="$"):
    if type(a) != type(b) and not (isinstance(a, (int, float)) and isinstance(b, (int, float))):
        return f"{path}: type {type(a).__name__} vs {type(b).__name__}"
    if isinstance(a, dict):
        if set(a) != set(b):
            return f"{path}: keys {sorted(set(a) ^ set(b))}"
        for k in a:
            d = first_difference(a[k], b[k], f"{path}.{k}")
            if d:
                return d
    elif isinstance(a, list):
        if len(a) != len(b):
            return f"{path}: length {len(a)} vs {len(b)}"
        for i, (x, y) in enumerate(zip(a, b)):
            d = first_difference(x, y, f"{path}[{i}]")
            if d:
                return d
    elif a != b:
        return f"{path}: {a!r} vs {b!r}"
    return None


SRC = INDEX.read_text(encoding="utf-8")


class PageBlocks(unittest.TestCase):
    def test_block_markers_once(self):
        marks = ["/* SESSION23-DATA-START", "/* SESSION23-DATA-END", "/* SESSION23-ENGINE-START", "/* SESSION23-ENGINE-END"]
        pos = []
        for m in marks:
            self.assertEqual(SRC.count(m), 1, m)
            pos.append(SRC.index(m))
        self.assertEqual(pos, sorted(pos))
        self.assertLess(pos[-1], SRC.index("function renderPredictive(){"))
        self.assertGreater(pos[0], SRC.index("/* ASHTAKAVARGA-END */"))

    def test_page_data_block_is_current(self):
        import build_session23 as b23
        a = SRC.index(b23.DATA_START)
        z = SRC.index(b23.DATA_END) + len(b23.DATA_END) + 1
        self.assertEqual(SRC[a:z], b23.page_data_block(json.loads(JSON_PATH.read_text(encoding="utf-8"))),
                         "run python3 build_session23.py to refresh the page's data block")

    def test_no_duplicate_function_declarations_added(self):
        names = re.findall(r"^\s*function\s+([A-Za-z0-9_$]+)\s*\(", SRC, re.M)
        dup = sorted({n for n in names if names.count(n) > 1})
        # two renderClassificationSection declarations pre-date this work
        self.assertEqual([n for n in dup if n != "renderClassificationSection"], [])
        engine = SRC[SRC.index("/* SESSION23-ENGINE-START"):SRC.index("/* SESSION23-ENGINE-END")]
        mine = re.findall(r"function\s+([A-Za-z0-9_$]+)\s*\(", engine)
        self.assertTrue(mine and all(n.startswith("s23") for n in mine), mine)
        before = SRC[:SRC.index("/* SESSION23-ENGINE-START")] + SRC[SRC.index("/* SESSION23-ENGINE-END"):]
        self.assertEqual([n for n in mine if re.search(r"function\s+" + n + r"\s*\(", before)], [])


@unittest.skipUnless(PW, "playwright + chromium needed")
class JsEngine(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import tempfile
        cls.tmp_dir = tempfile.TemporaryDirectory()
        cls.tmp = pathlib.Path(cls.tmp_dir.name)

    @classmethod
    def tearDownClass(cls):
        cls.tmp_dir.cleanup()

    def test_generated_data_equals_json(self):
        rules = json.loads(JSON_PATH.read_text(encoding="utf-8"))
        want = {k: v for k, v in rules.items() if k not in ("reference", "meta")}
        got = run_page("return S23;", None, self.tmp)
        self.assertIsNone(first_difference(want, got))

    def test_js_equals_python_on_2000_random_charts(self):
        charts = random_charts(2000, 2310)
        body = JS_CANON + "return job.map(c => { const r = layers(c); const o = {}; for (const k in r) o[k] = fp(r[k]); return o; });"
        got = run_page(body, charts, self.tmp)
        bad = []
        for i, (c, g) in enumerate(zip(charts, got)):
            mine = python_layers(c)
            for k, v in mine.items():
                if fingerprint(v) != g[k]:
                    bad.append((i, k))
        if bad:
            i, k = bad[0]
            full = run_page(JS_CANON + "return layers(job);", charts[i], self.tmp)[k]
            self.fail(f"{len(bad)} mismatches; first: chart {i} layer {k}: "
                      f"{first_difference(python_layers(charts[i])[k], full)}")

    def test_partial_chart_new_layers_equal_python(self):
        # a chart whose Moon or nodes are not known (no birth time for the Moon's sign, a missing entry …)
        charts = random_charts(300, 4242)
        for i, c in enumerate(charts):
            for p in (("Rahu", "Ketu"), ("Moon",), ("Moon", "Rahu", "Ketu"))[i % 3]:
                c["signs"].pop(p)
        body = JS_CANON + """return job.map(c => {
            const lg = c.lagna + 1, pbr = {};
            for (const p of PL) if (c.signs[p] !== undefined) pbr[p] = c.signs[p] + 1;
            const ctx = ctxOf(c, lg, pbr);
            return {bhava_bhava: fp(s23BhavaBhava(ctx)), conditions: fp(s23ChartConditions(ctx)), life_areas: fp(s23LifeAreas(ctx))};
        });"""
        got = run_page(body, charts, self.tmp)
        bad = []
        for i, (c, g) in enumerate(zip(charts, got)):
            ctx = ctx_of(c)
            mine = {"bhava_bhava": br.bhava_bhava(RULES, ctx), "conditions": br.chart_conditions(RULES, ctx),
                    "life_areas": br.life_areas(RULES, ctx)}
            bad += [(i, k) for k, v in mine.items() if fingerprint(v) != g[k]]
        self.assertEqual(bad, [])

    def test_moon_waxing_matches_python_and_the_page_tithi(self):
        rnd = random.Random(7)
        pairs = [[rnd.uniform(0, 360), rnd.uniform(0, 360)] for _ in range(3000)]
        pairs += [[10.0, 10.0], [0.0, 180.0], [0.0, 179.999], [100.0, 99.0], [350.0, 170.0], [350.0, 169.9]]
        body = "return job.map(([s, m]) => [s23MoonWaxing(s, m), computeTithi(s, m).tithiNum <= 15]);"
        got = run_page(body, pairs, self.tmp)
        for (s, m), (js, tithi) in zip(pairs, got):
            self.assertEqual((js, tithi), (br.moon_waxing(s, m),) * 2, (s, m))

    def test_oracle_chart_in_the_page(self):
        c = {"lagna": 4, "signs": {"Moon": 6, "Ketu": 7, "Saturn": 9, "Jupiter": 10, "Mars": 3, "Sun": 2,
                                   "Venus": 2, "Mercury": 1, "Rahu": 1},
             "waxing": True, "degs": {p: 12.5 for p in br.PLANET_ORDER}}
        got = run_page(JS_CANON + "const r = layers(job); return {roles: r.roles, asp: r.graha_graha.aspects.map(a => [a.by, a.to, a.house_aspect])};",
                       c, self.tmp)
        self.assertEqual(got["roles"]["Saturn"], ["Maraka lord", "Dusthana lord", "Trishadaya lord"])
        self.assertEqual(len(got["asp"]), 12)


BROWSER_SCRIPT = r"""
const { chromium } = require(process.env.PW_MODULE);
const indexUrl = 'file://' + process.argv[2];
const shots = process.env.S23_ARTIFACT_DIR || '';
const BIRTHS = [
  {dob:'1990-05-15', tob:'06:30:00', lat:'19.08',  lon:'72.88',  tz:'5.5'},
  {dob:'1975-11-02', tob:'22:10:00', lat:'13.0',   lon:'77.6',   tz:'5.5'},
  {dob:'2001-02-28', tob:'13:45:00', lat:'51.5',   lon:'-0.12',  tz:'0'},
];
const EXTRACT = () => {
  const $ = (s, r = document) => [...r.querySelectorAll(s)];
  const norm = s => s.replace(/\s+/g, ' ').trim();
  const classes = {};
  $('#s23-classes-body table.s23-table[data-class]').forEach(t => {
    classes[t.dataset.class] = $('tbody tr[data-house]', t).map(tr => ({
      house: +tr.dataset.house, sign: +tr.dataset.sign, lord: tr.dataset.lord,
      lord_house: tr.dataset.lordHouse === '' ? null : +tr.dataset.lordHouse,
      skt: tr.querySelector('.s23-skt').textContent, abbr: tr.querySelector('.s23-abbr').textContent,
      planets: $('.s23-pl', tr).map(x => x.dataset.p), text: norm(tr.innerText)}));
  });
  const m = document.getElementById('s23-matrix');
  const bad = document.getElementById('s23-badhaka');
  const predict = document.getElementById('s23-predict-body');
  return {
    sid: currentChart.chart.sidereal, lagna: currentChart.lagnaRashi.n,
    classes, titles: $('#report details.section .section-title').map(x => x.textContent),
    classHeadings: $('#s23-classes-body .s23-class-block .s23-class-head').map(x => norm(x.innerText)),
    matrixHead: m ? $('thead th[data-class]', m).map(th => th.dataset.class) : null,
    matrix: m ? $('tbody tr[data-house]', m).map(tr => ({house: +tr.dataset.house,
              classes: $('td[data-class]', tr).filter(td => td.dataset.in === '1').map(td => td.dataset.class),
              cells: $('td[data-class]', tr).length})) : null,
    badhaka: bad ? {mode: bad.dataset.mode, house: +bad.dataset.house, sign: +bad.dataset.sign, lord: bad.dataset.lord,
                    occupants: bad.dataset.occupants ? bad.dataset.occupants.split(',') : [], text: norm(bad.innerText)} : null,
    readings: $('#s23-readings tbody tr[data-planet]').map(tr => ({planet: tr.dataset.planet, house: +tr.dataset.house,
              nature: tr.dataset.nature, classes: tr.dataset.classes ? tr.dataset.classes.split(',') : [],
              texts: $('li', tr).map(li => norm(li.innerText))})),
    fivestep: $('#s23-fivestep tbody tr[data-house]').map(tr => ({house: +tr.dataset.house, lord: tr.dataset.lord,
              lord_house: +tr.dataset.lordHouse, occupants: tr.dataset.occupants ? tr.dataset.occupants.split(',') : [],
              aspecting: tr.dataset.aspecting ? tr.dataset.aspecting.split(',') : [], karakas: tr.dataset.karakas ? tr.dataset.karakas.split(',') : []})),
    gb: $('#s23-predict-body .s23-gb').map(c => ({planet: c.dataset.planet, house: +c.dataset.house, status: c.dataset.status,
              dig: norm(c.querySelector('.s23-dig').textContent), digLine: norm((c.querySelector('.s23-dig-line') || {textContent: ''}).textContent),
              digbala: norm((c.querySelector('.s23-digbala-line') || {textContent: ''}).textContent),
              points: $('.s23-points > li', c).map(li => norm(li.innerText)),
              extra: $('.s23-extra > li', c).map(li => norm(li.innerText)),
              classTexts: $('.s23-class-texts > li', c).map(li => norm(li.innerText)),
              tags: $('.s23-tag', c).map(x => x.textContent)})),
    bb: $('#s23-bb tbody tr[data-lord-of]').map(tr => ({lord_of: +tr.dataset.lordOf, sits_in: +tr.dataset.sits, lord: tr.dataset.lord,
              status: tr.dataset.status, text: norm(tr.querySelector('.s23-text').textContent), tags: $('.s23-tag', tr).map(x => x.textContent),
              exchange: +(tr.dataset.exchange || 0), dictum: tr.dataset.dictum === '1',
              signLine: norm(tr.querySelector('.s23-sign-line').textContent), placement: norm(tr.querySelector('.s23-placement').textContent)})),
    gr: $('#s23-gr tbody tr[data-planet]').map(tr => ({planet: tr.dataset.planet, sign: +tr.dataset.sign, status: tr.dataset.status,
              tatwa: tr.dataset.tatwa, direction: tr.dataset.direction, varna: tr.dataset.varna, mode: tr.dataset.mode,
              dignity: tr.dataset.dignity, text: norm(tr.nextElementSibling.textContent), tags: $('.s23-tag', tr).map(x => x.textContent)})),
    conj: $('#s23-conj > li').map(li => ({a: li.dataset.a, b: li.dataset.b, house: +li.dataset.house, status: li.dataset.status,
              combust: li.dataset.combust, text: norm(li.querySelector('.s23-text').textContent),
              note: norm((li.querySelector('.s23-combust') || {textContent: ''}).textContent)})),
    asp: $('#s23-asp > li').map(li => ({by: li.dataset.by, to: li.dataset.to, h: +li.dataset.h, status: li.dataset.status,
              text: norm(li.querySelector('.s23-text').textContent), jup: !!li.querySelector('.s23-jup'),
              meanings: $('.s23-meanings li', li).map(x => norm(x.textContent))})),
    predictText: predict ? predict.innerText : '',
    classesText: document.getElementById('s23-classes-body').innerText,
  };
};
(async () => {
  const browser = await chromium.launch({headless: true});
  const ctx = await browser.newContext({viewport: {width: 1280, height: 900}});
  const blocked = [];
  await ctx.route('**/*', r => r.request().url().startsWith('file:') ? r.continue()
                                : (blocked.push(r.request().url()), r.abort()));
  const page = await ctx.newPage();
  const jsErrors = [];
  page.on('pageerror', e => jsErrors.push(String(e)));
  await page.goto(indexUrl);
  await page.click('#btn-manual');
  const out = {charts: [], jsErrors, blocked: null};
  for (const [i, b] of BIRTHS.entries()) {
    await page.fill('#f-dob', b.dob);
    await page.fill('#f-tob', b.tob);
    await page.fill('#f-lat', b.lat); await page.press('#f-lat', 'Tab');
    await page.fill('#f-lon', b.lon); await page.press('#f-lon', 'Tab');
    await page.fill('#f-tz', b.tz);
    await page.click('#btn-generate');
    await page.waitForSelector('#report.show');
    await page.click('#reportNav button:has-text("Bhāvas")');
    const d = await page.evaluate(EXTRACT);
    d.nav = await page.evaluate(() => {
      const grp = [...document.querySelectorAll('#reportNav button')].map(b => b.textContent);
      const mine = [...document.querySelectorAll('#report details.section')].filter(s => /Bhāva Analysis/.test(s.querySelector('.section-title').textContent))
        .map(s => ({title: s.querySelector('.section-title').textContent, grp: s.dataset.grp, shown: getComputedStyle(s).display !== 'none'}));
      return {grp, mine};
    });
    out.charts.push(d);
    if (i === 0) {
      await page.click('#reportNav button:has-text("Planetary")');
      out.hiddenElsewhere = await page.evaluate(() => [...document.querySelectorAll('#report details.section')]
        .filter(s => /Bhāva Analysis/.test(s.querySelector('.section-title').textContent)).map(s => getComputedStyle(s).display));
      await page.click('#reportNav button:has-text("Bhāvas")');
      out.crowded = await page.evaluate(() => {
        // every planet in one rāśi (Karka, 100°) with Rahu/Ketu in the same sign: all tables must still render
        const sid = {}; for (const p of S23_PLANETS) sid[p] = 100 + S23_PLANETS.indexOf(p) * 0.1;
        renderS23Sections(sid, 4);     // Lagna Karka: every planet in the 1st house
        const q = document.querySelector('#s23-classes-body table[data-class="Kendra"]');
        return {tables: document.querySelectorAll('#s23-classes-body table, #s23-predict-body table').length,
                kendraPlanets: [...q.querySelectorAll('.s23-pl')].map(x => x.dataset.p),
                cards: document.querySelectorAll('#s23-predict-body .s23-gb').length,
                conj: document.querySelectorAll('#s23-conj > li[data-a]').length,
                text: document.getElementById('s23-predict-body').innerText.length}; });
      await page.evaluate(() => renderS23Sections(currentChart.chart.sidereal, currentChart.lagnaRashi.n));
      out.noLagna = await page.evaluate(() => { renderS23Sections(currentChart.chart.sidereal, 0);
        return {classes: document.getElementById('s23-classes-body').innerText, predict: document.getElementById('s23-predict-body').innerText,
                tables: document.querySelectorAll('#s23-classes-body table, #s23-predict-body table').length}; });
      await page.evaluate(() => renderS23Sections(currentChart.chart.sidereal, currentChart.lagnaRashi.n));
      out.layout = {};
      if (shots) await page.locator('#s23-classes-body').screenshot({path: shots + '/s23_classes_desktop.png'});
      if (shots) await page.locator('#s23-predict-body').screenshot({path: shots + '/s23_predict_desktop.png'});
      await page.setViewportSize({width: 375, height: 800});
      out.layout.phone = await page.evaluate(() => ({vw: window.innerWidth, pageScrollW: document.documentElement.scrollWidth,
        wrappers: [...document.querySelectorAll('#s23-classes-body .table-scroll, #s23-predict-body .table-scroll')].map(w => {
          const tb = w.querySelector('table');
          return {w: Math.round(w.getBoundingClientRect().width), id: tb.id || tb.dataset.class || '',
                  scrolls: w.scrollWidth > w.clientWidth || tb.scrollWidth > tb.clientWidth}; })}));
      if (shots) await page.locator('#s23-classes-body').screenshot({path: shots + '/s23_classes_phone.png'});
      await page.setViewportSize({width: 1280, height: 900});
      await page.emulateMedia({media: 'print'});
      out.layout.print = await page.evaluate(() => {
        const sections = [...document.querySelectorAll('#report details.section')].filter(s => /Bhāva Analysis/.test(s.querySelector('.section-title').textContent));
        const skt = document.querySelector('#s23-classes-body .s23-skt'), abbr = document.querySelector('#s23-classes-body .s23-abbr');
        return {display: sections.map(s => getComputedStyle(s).display), skt: getComputedStyle(skt).display, abbr: getComputedStyle(abbr).display};
      });
      await page.setViewportSize({width: 695, height: 1000});
      out.layout.printA4 = await page.evaluate(() =>
        [...document.querySelectorAll('#s23-classes-body .table-scroll, #s23-predict-body .table-scroll')].map(w => {
          const t = w.querySelector('table');
          return {id: t.id || t.dataset.class || '', wrapClient: w.clientWidth, wrapScroll: w.scrollWidth, tableScroll: t.scrollWidth,
                  headers: [...t.querySelectorAll('thead th')].map(th => ({txt: th.innerText.replace(/\s+/g, ' ').trim(), cw: th.clientWidth,
                    sw: (() => { const r = document.createRange(); r.selectNodeContents(th); return Math.ceil(r.getBoundingClientRect().width); })()}))};
        }));
      await page.setViewportSize({width: 1280, height: 900});
      if (shots) await page.pdf({path: shots + '/s23_print.pdf', format: 'A4', printBackground: true});
      await page.emulateMedia({media: 'screen'});
    }
  }
  out.blocked = blocked;
  console.log(JSON.stringify(out));
  await browser.close();
})().catch(e => { console.error(e); process.exit(1); });
"""


def _norm(s):
    return " ".join(str(s).split())


def page_chart(d):
    sid = d["sid"]
    signs = {p: int((sid[p] % 360) // 30) for p in br.PLANET_ORDER}
    degs = {p: (sid[p] % 360) - 30 * signs[p] for p in br.PLANET_ORDER}
    return d["lagna"] - 1, signs, degs, br.moon_waxing(sid["Sun"], sid["Moon"])


@unittest.skipUnless(PW, "playwright + chromium needed")
class BrowserSession23(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        env = dict(os.environ, PW_MODULE=PW, PLAYWRIGHT_BROWSERS_PATH="/opt/pw-browsers")
        import tempfile
        cls.tmp_dir = tempfile.TemporaryDirectory()
        script = pathlib.Path(cls.tmp_dir.name) / "browser.js"
        script.write_text(BROWSER_SCRIPT, encoding="utf-8")
        p = subprocess.run([NODE, str(script), str(INDEX)], capture_output=True, text=True, timeout=600, env=env)
        if p.returncode != 0:
            raise AssertionError(p.stderr[-2000:])
        cls.out = json.loads(p.stdout.strip().splitlines()[-1])

    @classmethod
    def tearDownClass(cls):
        cls.tmp_dir.cleanup()

    def test_no_js_errors_and_no_network(self):
        self.assertEqual(self.out["jsErrors"], [])
        # the page's own font and sheet-logger requests are blocked; the new sections must not add any host
        import urllib.parse
        hosts = {urllib.parse.urlparse(u).netloc for u in self.out["blocked"]}
        self.assertLessEqual(hosts, {"fonts.googleapis.com", "script.google.com"})

    def test_sections_sit_in_the_bhavas_group(self):
        for d in self.out["charts"]:
            self.assertEqual([m["title"] for m in d["nav"]["mine"]],
                             ["Bhāva Analysis — the 12 Houses", "Bhāva Analysis — The 10 Bhāva Classes",
                              "Bhāva Analysis — Prediction Layers"])
            self.assertTrue(all(m["shown"] and m["grp"] == "3" for m in d["nav"]["mine"]), d["nav"])
        self.assertEqual(self.out["hiddenElsewhere"], ["none", "none", "none"])

    def test_ten_class_tables_equal_python(self):
        for d in self.out["charts"]:
            lg, signs, degs, w = page_chart(d)
            want = br.classification_tables(RULES, lg, signs)
            self.assertEqual(list(d["classes"]), br.CLASS_ORDER)
            for name, rows in want.items():
                got = d["classes"][name]
                self.assertEqual([(r["house"], r["sign"], r["planets"], r["lord"], r["lord_house"]) for r in got],
                                 [(r["house"], r["sign"], r["planets"], r["lord"], r["lord_house"]) for r in rows], name)
                for r in got:
                    self.assertEqual(r["skt"], RULES["reference"]["RASHI"][r["sign"]]["sanskrit"])
                    self.assertEqual(r["abbr"], RULES["reference"]["RASHI"][r["sign"]]["english"][:3])
            self.assertEqual(len(d["classHeadings"]), 10)
            self.assertIn("Apachaya", d["classHeadings"][5])
            self.assertIn("1, 2, 4, 7, 8", d["classHeadings"][5])

    def test_matrix_has_12_rows_10_classes(self):
        for d in self.out["charts"]:
            lg, *_ = page_chart(d)
            self.assertEqual(d["matrixHead"], br.CLASS_ORDER)
            self.assertEqual(len(d["matrix"]), 12)
            self.assertTrue(all(r["cells"] == 10 for r in d["matrix"]))
            self.assertEqual([(r["house"], r["classes"]) for r in d["matrix"]],
                             [(r["house"], r["classes"]) for r in br.house_class_matrix(RULES, lg)])

    def test_badhaka_block(self):
        for d in self.out["charts"]:
            lg, signs, degs, w = page_chart(d)
            want = br.badhaka(RULES, lg, signs)
            b = d["badhaka"]
            self.assertEqual((b["mode"], b["house"], b["sign"], b["lord"], b["occupants"]),
                             (want["mode"], want["house"], want["sign"], want["lord"], want["occupants"]))
            self.assertIn("a sign, not the horoscope", b["text"])

    def test_readings_and_five_step_equal_python(self):
        for d in self.out["charts"]:
            lg, signs, degs, w = page_chart(d)
            want = br.class_readings(RULES, lg, signs)
            self.assertEqual([(r["planet"], r["house"], r["classes"], [_norm(x) for x in r["texts"]]) for r in want],
                             [(r["planet"], r["house"], r["classes"], r["texts"]) for r in d["readings"]])
            for r in d["readings"]:
                self.assertEqual(r["nature"], br.nature(r["planet"]))
            for h, got in zip(range(1, 13), d["fivestep"]):
                f = br.five_step(RULES, lg, signs, h)
                self.assertEqual((got["house"], got["lord"], got["lord_house"], got["occupants"]),
                                 (h, f["lord"], f["lord_house"], f["occupants"]))
                self.assertEqual(got["aspecting"], [f"{a['by']}:{a['house_aspect']}" for a in f["aspecting"]])
                self.assertEqual(got["karakas"], [f"{k['planet']}:{k['house']}" for k in f["karakas"]])

    def test_prediction_layers_equal_python_and_show_status_tags(self):
        for d in self.out["charts"]:
            lg, signs, degs, w = page_chart(d)
            for g, p in zip(d["gb"], br.PLANET_ORDER):
                want = br.graha_bhava(RULES, lg, signs, p, degs[p], w)
                self.assertEqual((g["planet"], g["house"], g["status"]), (p, want["house"], want["status"]))
                self.assertEqual(g["dig"], want["dignity"]["label"])
                self.assertEqual(g["digLine"], _norm(want["dignity_line"]))
                self.assertEqual(g["digbala"], _norm(want["digbala_line"]))
                self.assertEqual(g["points"], [_norm(x) for x in want["points"]])
                self.assertEqual(g["extra"], [_norm(x) for x in want["extra"]])
                self.assertEqual(g["classTexts"], [_norm(x) for x in want["class_texts"]])
                self.assertIn(want["status"], g["tags"])
            bb = br.bhava_bhava(RULES, br.context(lg, signs, degs))
            self.assertEqual(len(d["bb"]), len(bb))
            for r, want in zip(d["bb"], bb):
                self.assertEqual((r["lord_of"], r["sits_in"], r["lord"], r["status"], r["text"], r["exchange"], r["dictum"]),
                                 (want["lord_of"], want["sits_in"], want["lord"], want["status"], _norm(want["text"]),
                                  want["exchange"] or 0, bool(want["dictum"])))
                self.assertEqual(r["signLine"], _norm(want["sign_line"]))
                self.assertEqual(r["placement"], _norm(want["placement_line"]))
                self.assertIn(want["status"], r["tags"])
            for r, want in zip(d["gr"], br.graha_rashi(RULES, lg, signs, degs)):
                self.assertEqual((r["planet"], r["sign"], r["status"], r["tatwa"], r["direction"], r["varna"], r["mode"], r["dignity"], r["text"]),
                                 (want["planet"], want["sign"], want["status"], want["tatwa"], want["direction"], want["varna"],
                                  want["mode"], want["dignity"], _norm(want["text"])))
            gg = br.graha_graha(RULES, lg, signs, degs)
            self.assertEqual([(c["a"], c["b"], c["house"], c["status"], _norm(c["text"])) for c in gg["conjunctions"]],
                             [(c["a"], c["b"], c["house"], c["status"], c["text"]) for c in d["conj"]])
            self.assertEqual([c["note"] for c in d["conj"]], [_norm(c["combust_note"]) for c in gg["conjunctions"]])
            self.assertEqual([(a["by"], a["to"], a["house_aspect"], a["status"], _norm(a["text"]), a["jupiter_flag"],
                              [_norm(m) for m in a["meanings"]]) for a in gg["aspects"]],
                             [(a["by"], a["to"], a["h"], a["status"], a["text"], a["jup"], a["meanings"]) for a in d["asp"]])
            self.assertIn("taught", d["predictText"].lower())      # the legend explains the tags
            self.assertIn("curated", d["predictText"].lower())

    def test_crowded_chart_still_renders(self):
        c = self.out["crowded"]
        self.assertEqual(c["tables"], 15)                      # matrix, ten classes, readings, five-step, bhava+bhava, graha+rashi
        self.assertEqual(c["cards"], 9)
        self.assertEqual(c["conj"], 36)                      # nine planets in one sign: every pair is a conjunction
        self.assertEqual(c["kendraPlanets"], br.PLANET_ORDER)  # Lagna 4 is Karka, so all nine sit in the 1st, a Kendra
        self.assertGreater(c["text"], 5000)

    def test_sections_explain_missing_lagna(self):
        n = self.out["noLagna"]
        self.assertEqual(n["tables"], 0)
        self.assertIn("Lagna", n["classes"])
        self.assertIn("Lagna", n["predict"])

    def test_phone_tables_scroll_in_their_box(self):
        ph = self.out["layout"]["phone"]
        self.assertLessEqual(ph["pageScrollW"], ph["vw"] + 1, "the page itself must not scroll sideways")
        self.assertTrue(ph["wrappers"])
        for w in ph["wrappers"]:
            self.assertLessEqual(w["w"], ph["vw"], w)
        self.assertTrue(next(w for w in ph["wrappers"] if w["id"] == "s23-matrix")["scrolls"])

    def test_print_fits_a4_without_clipped_headers(self):
        pr = self.out["layout"]["print"]
        self.assertNotIn("none", pr["display"])
        self.assertEqual((pr["skt"], pr["abbr"]), ("none", "inline"))
        wraps = self.out["layout"]["printA4"]
        self.assertTrue(wraps)
        for w in wraps:
            self.assertLessEqual(w["wrapScroll"], w["wrapClient"] + 1, w)
            self.assertLessEqual(w["tableScroll"], w["wrapClient"] + 1, w)
            for h in w["headers"]:
                self.assertLessEqual(h["sw"], h["cw"] + 1, (w["id"], h))


DASHA_SCRIPT = r"""
const { chromium } = require(process.env.PW_MODULE);
const indexUrl = 'file://' + process.argv[2];
const shots = process.env.S23_ARTIFACT_DIR || '';
const NOW = '2026-10-08T12:00:00Z';
const BIRTHS = [
  {dob:'1990-05-15', tob:'06:30:00', lat:'19.08',  lon:'72.88',  tz:'5.5'},
  {dob:'1975-11-02', tob:'22:10:00', lat:'13.0',   lon:'77.6',   tz:'5.5'},
  {dob:'2001-02-28', tob:'13:45:00', lat:'51.5',   lon:'-0.12',  tz:'0'},
];
const EXTRACT = () => {
  const $ = (s, r = document) => [...r.querySelectorAll(s)];
  const norm = s => s.replace(/\s+/g, ' ').trim();
  const body = document.getElementById('s23-dasha-body');
  const run = body.querySelector('.s23-running');
  const part = k => { const e = run && run.querySelector('.s23-run-' + k);
    return e ? {lord: e.dataset.lord, roles: $('.s23-role', e).map(c => c.dataset.role), text: $('.s23-run-text li', e).map(li => norm(li.innerText))} : null; };
  return {
    sid: currentChart.chart.sidereal, lagna: currentChart.lagnaRashi.n, dob: currentChart.dob, tob: currentChart.tob,
    nav: $('#reportNav button').map(b => b.textContent),
    sectionTitles: $('#report details.section .section-title').map(x => x.textContent),
    dashaSection: (() => { const s = body.closest('details'); return {grp: s.dataset.grp, shown: getComputedStyle(s).display !== 'none', title: s.querySelector('.section-title').textContent}; })(),
    predictiveTabs: $('#reportTabs .rtab').map(b => b.dataset.tab), activeTab: $('#reportTabs .rtab.active').map(b => b.dataset.tab),
    panes: $('#report .rpane').map(p => ({id: p.id, display: getComputedStyle(p).display})),
    oldPane: !!document.getElementById('pane-dasha'),
    heads: $('#report h4.pane-h').filter(h => /Vimśottarī Daśā/.test(h.textContent)).length,
    running: run ? {maha: part('maha'), bhukti: part('bhukti'), dataMaha: run.dataset.maha, dataBhukti: run.dataset.bhukti} : null,
    roles: $('#s23-roles tbody tr[data-planet]').map(tr => ({planet: tr.dataset.planet, roles: tr.dataset.roles ? tr.dataset.roles.split(',') : [],
              owns: tr.dataset.owns ? tr.dataset.owns.split(',').map(Number) : [], in: +tr.dataset.in})),
    watch: $('#s23-watch tbody tr[data-maha]').map(tr => ({maha: tr.dataset.maha, bhukti: tr.dataset.bhukti, kind: tr.dataset.kind,
              start: +tr.dataset.start, end: +tr.dataset.end, roles: tr.dataset.roles.split(',')})),
    blocks: $('#s23-dasha-body .maha-block').map(b => ({maha: b.dataset.maha, headRoles: $('.maha-head .s23-role', b).map(c => c.dataset.role),
              rows: $('tbody tr[data-bhukti]', b).map(tr => ({bhukti: tr.dataset.bhukti, now: tr.classList.contains('now'),
                                                              roles: $('.s23-role', tr).map(c => c.dataset.role)}))})),
    text: body.innerText,
  };
};
(async () => {
  const browser = await chromium.launch({headless: true});
  const ctx = await browser.newContext({viewport: {width: 1280, height: 900}, timezoneId: 'UTC'});
  const blocked = [];
  await ctx.route('**/*', r => r.request().url().startsWith('file:') ? r.continue()
                                : (blocked.push(r.request().url()), r.abort()));
  const page = await ctx.newPage();
  const jsErrors = [];
  page.on('pageerror', e => jsErrors.push(String(e)));
  await page.clock.install({time: new Date(NOW)});
  await page.goto(indexUrl);
  await page.click('#btn-manual');
  const out = {charts: [], jsErrors, now: NOW};
  for (const [i, b] of BIRTHS.entries()) {
    await page.fill('#f-dob', b.dob);
    await page.fill('#f-tob', b.tob);
    await page.fill('#f-lat', b.lat); await page.press('#f-lat', 'Tab');
    await page.fill('#f-lon', b.lon); await page.press('#f-lon', 'Tab');
    await page.fill('#f-tz', b.tz);
    await page.click('#btn-generate');
    await page.waitForSelector('#report.show');
    await page.click('#reportNav button:has-text("Daśā")');
    const d = await page.evaluate(EXTRACT);
    d.tz = parseFloat(b.tz);
    out.charts.push(d);
    if (i === 0) {
      if (shots) await page.locator('#s23-dasha-body').screenshot({path: shots + '/s23_dasha_desktop.png'});
      out.predictiveHiddenOnDasha = await page.evaluate(() => [...document.querySelectorAll('#report details.section')]
        .filter(s => /Predictive/.test(s.querySelector('.section-title').textContent)).map(s => getComputedStyle(s).display));
      await page.click('#reportNav button:has-text("Predictive")');
      out.predictiveShown = await page.evaluate(() => ({
        dasha: getComputedStyle(document.getElementById('s23-dasha-body').closest('details')).display,
        transit: getComputedStyle(document.getElementById('pane-transit')).display,
        active: [...document.querySelectorAll('#reportTabs .rtab.active')].map(b => b.dataset.tab)}));
      await page.emulateMedia({media: 'print'});
      out.print = await page.evaluate(() => ({dasha: getComputedStyle(document.getElementById('s23-dasha-body').closest('details')).display,
        heads: [...document.querySelectorAll('#report h4.pane-h')].filter(h => /Vimśottarī Daśā/.test(h.textContent)).length}));
      await page.setViewportSize({width: 375, height: 800});
      out.phone = await page.evaluate(() => ({vw: window.innerWidth, pageScrollW: document.documentElement.scrollWidth}));
      await page.setViewportSize({width: 1280, height: 900});
      if (shots) await page.pdf({path: shots + '/s23_dasha_print.pdf', format: 'A4', printBackground: true});
      await page.emulateMedia({media: 'screen'});
      out.noLagna = await page.evaluate(() => { renderS23Dasha(buildDasha(currentChart.chart.sidereal.Moon, currentChart.dob, currentChart.tob), null, new Date());
        return document.getElementById('s23-dasha-body').innerText; });
      await page.evaluate(() => renderPredictive());
      out.runPart = await page.evaluate(() => ({
        onlyMahaRole: document.createElement('div').appendChild(Object.assign(document.createElement('div'), {innerHTML: s23RunPart("bhukti", "Bhukti", {lord: "Venus", roles: [], text: []}, ["Trishadaya lord"])})).innerText,
        noRole: document.createElement('div').appendChild(Object.assign(document.createElement('div'), {innerHTML: s23RunPart("bhukti", "Bhukti", {lord: "Sun", roles: [], text: []}, [])})).innerText}));
    }
  }
  out.blocked = blocked;
  console.log(JSON.stringify(out));
  await browser.close();
})().catch(e => { console.error(e); process.exit(1); });
"""


def _ms_naive(d):
    return d.replace(tzinfo=__import__("datetime").timezone.utc).timestamp() * 1000


@unittest.skipUnless(PW, "playwright + chromium needed")
class BrowserDasha(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        env = dict(os.environ, PW_MODULE=PW, PLAYWRIGHT_BROWSERS_PATH="/opt/pw-browsers")
        import tempfile
        cls.tmp_dir = tempfile.TemporaryDirectory()
        script = pathlib.Path(cls.tmp_dir.name) / "dasha.js"
        script.write_text(DASHA_SCRIPT, encoding="utf-8")
        p = subprocess.run([NODE, str(script), str(INDEX)], capture_output=True, text=True, timeout=600, env=env)
        if p.returncode != 0:
            raise AssertionError(p.stderr[-2000:])
        cls.out = json.loads(p.stdout.strip().splitlines()[-1])
        import datetime
        cls.now = datetime.datetime(2026, 10, 8, 12, 0, 0)

    @classmethod
    def tearDownClass(cls):
        cls.tmp_dir.cleanup()

    def python_dasha(self, d):
        import datetime
        lg, signs, degs, w = page_chart(d)
        y, m, dd = map(int, d["dob"].split("-"))
        hh, mm, ss = (list(map(int, d["tob"].split(":"))) + [0, 0, 0])[:3]
        # the page dates the daśā from the birth instant in UT (local clock time minus the UTC offset)
        birth = datetime.datetime(y, m, dd, hh, mm, ss) - datetime.timedelta(hours=d["tz"])
        dasha = br.vimshottari(RULES, d["sid"]["Moon"], birth, self.now)
        return lg, signs, dasha, br.planet_roles(RULES, lg, signs)

    def test_no_js_errors(self):
        self.assertEqual(self.out["jsErrors"], [])

    def test_top_level_nav_has_dasha_button(self):
        for d in self.out["charts"]:
            self.assertIn("Daśā–Bhukti", d["nav"])
            self.assertEqual(d["nav"].index("Daśā–Bhukti") + 1, d["nav"].index("Predictive & Remedies"))
            self.assertEqual(d["dashaSection"]["title"], "Daśā–Bhukti — Vimśottarī")
            self.assertTrue(d["dashaSection"]["shown"])
        self.assertEqual(self.out["predictiveHiddenOnDasha"], ["none"])

    def test_pill_removed_and_dasha_printed_once(self):
        for d in self.out["charts"]:
            self.assertNotIn("dasha", d["predictiveTabs"])
            self.assertFalse(d["oldPane"])
            self.assertEqual(d["heads"], 1)
            self.assertEqual(d["predictiveTabs"][0], "transit")
            self.assertEqual(d["activeTab"], ["transit"])
        self.assertEqual(self.out["print"], {"dasha": "block", "heads": 1})
        ps = self.out["predictiveShown"]
        self.assertEqual((ps["dasha"], ps["transit"], ps["active"]), ("none", "block", ["transit"]))
        self.assertLessEqual(self.out["phone"]["pageScrollW"], self.out["phone"]["vw"] + 1)

    def test_running_pair_matches_python(self):
        for d in self.out["charts"]:
            lg, signs, dasha, roles = self.python_dasha(d)
            want = br.running_roles(RULES, dasha, roles)
            run = d["running"]
            self.assertEqual((run["dataMaha"], run["dataBhukti"]), (dasha["cur_maha"]["lord"], dasha["cur_bhukti"]["lord"]))
            for k in ("maha", "bhukti"):
                self.assertEqual(run[k]["lord"], want[k]["lord"])
                self.assertEqual(run[k]["roles"], want[k]["roles"])
                self.assertEqual(run[k]["text"], [_norm(x) for x in want[k]["text"]])

    def test_roles_table_equals_python(self):
        for d in self.out["charts"]:
            lg, signs, dasha, roles = self.python_dasha(d)
            self.assertEqual([r["planet"] for r in d["roles"]], br.PLANET_ORDER)
            for r in d["roles"]:
                self.assertEqual(r["roles"], roles[r["planet"]])
                self.assertEqual(r["owns"], [h for h in range(1, 13) if br.house_lord(RULES, lg, h) == r["planet"]])
                self.assertEqual(r["in"], br.house_of(lg, signs[r["planet"]]))

    def test_watch_periods_equal_python(self):
        for d in self.out["charts"]:
            lg, signs, dasha, roles = self.python_dasha(d)
            want = br.watch_periods(RULES, dasha, roles, self.now, years=10)
            self.assertEqual(len(d["watch"]), len(want))
            for g, w in zip(d["watch"], want):
                self.assertEqual((g["maha"], g["bhukti"], g["kind"], g["roles"]), (w["maha"], w["bhukti"], w["kind"], w["roles"]))
                self.assertAlmostEqual(g["start"], _ms_naive(w["start"]), delta=1000)
                self.assertAlmostEqual(g["end"], _ms_naive(w["end"]), delta=1000)

    def test_role_chips_on_bhukti_rows(self):
        for d in self.out["charts"]:
            lg, signs, dasha, roles = self.python_dasha(d)
            self.assertEqual([b["maha"] for b in d["blocks"]], [t["lord"] for t in dasha["timeline"]])
            for b, t in zip(d["blocks"], dasha["timeline"]):
                self.assertEqual(b["headRoles"], roles[t["lord"]])
                self.assertEqual([r["bhukti"] for r in b["rows"]], [x["lord"] for x in t["bhuktis"]])
                for r in b["rows"]:
                    self.assertEqual(r["roles"], br._applicable_roles(roles[r["bhukti"]], False))
            current = [(b["maha"], r["bhukti"]) for b in d["blocks"] for r in b["rows"] if r["now"]]
            self.assertEqual(current, [(dasha["cur_maha"]["lord"], dasha["cur_bhukti"]["lord"])])

    def test_indicative_wording_present(self):
        for d in self.out["charts"]:
            self.assertIn("indicative", d["text"].lower())
            self.assertNotRegex(d["text"].lower(), r"will die|certainly|definitely")
        for d in self.out["charts"]:
            self.assertIn("before the time of death is promised", d["text"])      # the role legend quotes the slide

    def test_running_bhukti_whose_only_role_is_mahadasha_only_is_explained(self):
        p = self.out["runPart"]
        self.assertIn("Trishadaya", p["onlyMahaRole"])
        self.assertIn("Mahādaśā only", p["onlyMahaRole"])
        self.assertNotIn("holds none of the six roles", p["onlyMahaRole"])
        self.assertIn("holds none of the six roles", p["noRole"])

    def test_missing_lagna_still_shows_the_daśā_with_a_note(self):
        self.assertIn("Lagna", self.out["noLagna"])
        self.assertIn("Mahādaśā", self.out["noLagna"])


# ------------------------------------------------------------------------------------------ Excel calculator
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

CALC_NOW = None


def _lo_recalc(paths, outdir):
    import tempfile
    profile = tempfile.mkdtemp(prefix="lo_profile_")
    try:
        subprocess.run([SOFFICE, f"-env:UserInstallation={pathlib.Path(profile).resolve().as_uri()}", "--headless",
                        "--convert-to", "xlsx", "--outdir", str(outdir), *map(str, paths)],
                       check=True, capture_output=True, timeout=900)
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def _waxing_of(c):
    lon = lambda p: c["signs"][p] * 30 + c["degs"][p]
    return br.moon_waxing(lon("Sun"), lon("Moon"))


def calc_charts(n, seed):
    import datetime
    rnd = random.Random(seed)
    out = []
    for c in random_charts(n, seed):
        c["birth"] = datetime.datetime(rnd.randint(1940, 2010), rnd.randint(1, 12), rnd.randint(1, 28),
                                       rnd.randint(0, 23), rnd.randint(0, 59), 0)
        c["waxing"] = _waxing_of(c)       # the sheet derives it from the Sun and Moon longitudes
        out.append(c)
    return out


def condition_charts():
    """The teacher's worked charts plus charts that make the rarer conditional readings hold."""
    import datetime
    import test_teacher_charts as tc
    out = []
    for name in ("S23_2024_A", "S23_2024_B", "S23_2024_C", "S23_2024_D", "S27_CHART_2"):
        c = dict(tc.CHARTS[name], birth=datetime.datetime(1995, 3, 10, 8, 30), gender="Female")
        out.append(c)
    base = {p: 7 for p in br.PLANET_ORDER}                                  # Simha Lagna; all in Vrischika (4th) …
    out.append({"lagna": 4, "signs": dict(base, Saturn=4, Rahu=1, Ketu=7), "degs": {p: 12.0 for p in br.PLANET_ORDER},
                "retro": {"Saturn": True, "Rahu": True, "Ketu": True}, "birth": datetime.datetime(2001, 6, 1, 4, 0), "gender": "Male"})
    out.append({"lagna": 4, "signs": dict(base, Saturn=1, Sun=1, Rahu=2, Ketu=8), "degs": dict({p: 12.0 for p in br.PLANET_ORDER}, Saturn=10.0),
                "retro": {}, "birth": datetime.datetime(2002, 2, 2, 2, 0), "gender": None})   # Saturn combust in the 10th, young
    for c in out:
        c["waxing"] = _waxing_of(c)
    return out


def dignity_boundary_charts(n=36):
    """Every planet placed at sign/degree values that straddle the exaltation, debilitation and MT boundaries."""
    import datetime
    ref = RULES["reference"]
    charts = []
    for k in range(n):
        signs, degs = {}, {}
        for i, p in enumerate(br.PLANET_ORDER):
            dd = ref["DIGNITY_DEG"].get(p)
            pool = [0.0, 0.5, 14.99, 29.99]
            if dd:
                pool += [dd["exD"], dd["exD"] - 1, dd["exD"] + 1, dd["exD"] - 1.01, dd["exD"] + 1.01,
                         dd["deD"], dd["deD"] - 1, dd["deD"] + 1, dd["mt"][0], dd["mt"][1], dd["mt"][1] + 0.01,
                         max(dd["mt"][0] - 0.01, 0)]
            degs[p] = round(min(pool[(k * 5 + i * 3) % len(pool)], 29.99), 2)     # a degree in a sign is below 30
            home = [dd["exR"], dd["deR"], dd["mtR"]][k % 3] - 1 if dd else (k + i) % 12
            signs[p] = home if (k // 3 + i) % 2 == 0 else (home + k // 3 + i) % 12
        c = {"lagna": k % 12, "signs": signs, "degs": degs, "birth": datetime.datetime(1990, 1, 1)}
        c["waxing"] = _waxing_of(c)
        charts.append(c)
    return charts


@unittest.skipUnless(SOFFICE and openpyxl, "LibreOffice and openpyxl are needed for the Excel check")
class ExcelSession23(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import datetime
        import tempfile
        import make_session23_xlsx as mx
        cls.mx = mx
        cls.tmp_dir = tempfile.TemporaryDirectory()
        tmp = pathlib.Path(cls.tmp_dir.name)
        cls.now = datetime.datetime(2026, 10, 8, 12, 0, 0)
        base = tmp / "base.xlsx"
        mx.build_workbook(base)
        cls.charts = calc_charts(24, 77) + dignity_boundary_charts() + condition_charts()
        src = tmp / "in"
        src.mkdir()
        paths = []
        for i, c in enumerate(cls.charts):
            wb = openpyxl.load_workbook(base)
            mx.set_inputs(wb, c, cls.now)
            f = src / f"chart{i:02d}.xlsx"
            wb.save(f)
            paths.append(f)
        out = tmp / "out"
        out.mkdir()
        _lo_recalc(paths, out)
        cls.books = [openpyxl.load_workbook(out / p.name, data_only=True) for p in paths]

    @classmethod
    def tearDownClass(cls):
        cls.tmp_dir.cleanup()

    def each(self, count=None):
        for i, (c, wb) in enumerate(zip(self.charts, self.books)):
            if count is None or i < count:
                yield i, c, wb

    def test_class_tables_equal_python(self):
        mx = self.mx
        for i, c, wb in self.each(24):
            ws = wb[mx.PREFIX + "Classes_Calc"]
            want = br.classification_tables(RULES, c["lagna"], c["signs"])
            for name in br.CLASS_ORDER:
                rows = [r for r in range(mx.CLASS_ROW0, mx.CLASS_ROW0 + mx.CLASS_ROWS) if ws.cell(row=r, column=1).value == name]
                got = [(ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value,
                        ws.cell(row=r, column=4).value or "", ws.cell(row=r, column=5).value, ws.cell(row=r, column=6).value) for r in rows]
                exp = [(r["house"], RULES["reference"]["RASHI"][r["sign"]]["sanskrit"], ", ".join(r["planets"]), r["lord"], r["lord_house"])
                       for r in want[name]]
                self.assertEqual(got, exp, (i, name))

    def test_matrix_and_badhaka_equal_python(self):
        mx = self.mx
        for i, c, wb in self.each(24):
            ws = wb[mx.PREFIX + "Classes_Calc"]
            want = br.house_class_matrix(RULES, c["lagna"])
            got = [[ws.cell(row=mx.MATRIX_ROW0 + 1 + h, column=3 + k).value for k in range(10)] for h in range(12)]
            self.assertEqual(got, [[1 if n in r["classes"] else 0 for n in br.CLASS_ORDER] for r in want], i)
            self.assertEqual([ws.cell(row=mx.MATRIX_ROW0 + 1 + h, column=13).value or "" for h in range(12)],
                             [", ".join(r["classes"]) for r in want], i)
            b = br.badhaka(RULES, c["lagna"], c["signs"])
            self.assertEqual((ws.cell(row=mx.BADHAKA_ROW, column=2).value, ws.cell(row=mx.BADHAKA_ROW, column=3).value,
                              ws.cell(row=mx.BADHAKA_ROW, column=4).value, ws.cell(row=mx.BADHAKA_ROW, column=5).value,
                              ws.cell(row=mx.BADHAKA_ROW, column=6).value or ""),
                             (b["mode"], b["house"], RULES["reference"]["RASHI"][b["sign"]]["sanskrit"], b["lord"], ", ".join(b["occupants"])), i)

    def test_roles_equal_python(self):
        mx = self.mx
        for i, c, wb in self.each(24):
            ws = wb[mx.PREFIX + "Roles_Calc"]
            want = br.planet_roles(RULES, c["lagna"], c["signs"])
            for k, p in enumerate(br.PLANET_ORDER):
                r = mx.ROLES_ROW0 + k
                self.assertEqual(ws.cell(row=r, column=1).value, p)
                flags = [ws.cell(row=r, column=2 + j).value for j in range(6)]
                self.assertEqual(flags, [1 if role in want[p] else 0 for role in br.ROLE_ORDER], (i, p))
                self.assertEqual(ws.cell(row=r, column=8).value or "", ", ".join(want[p]), (i, p))

    def test_graha_bhava_text_equals_python(self):
        mx = self.mx
        for i, c, wb in self.each(24):
            ws = wb[mx.PREFIX + "Predict_Calc"]
            for k, p in enumerate(br.PLANET_ORDER):
                r = mx.PREDICT_ROW0 + k
                want = br.graha_bhava(RULES, c["lagna"], c["signs"], p, c["degs"][p], c["waxing"])
                got = {name: ws.cell(row=r, column=col).value for name, col in mx.PREDICT_COLS.items()}
                self.assertEqual(got["planet"], p)
                self.assertEqual(got["house"], want["house"], (i, p))
                self.assertEqual(got["nature"], br.nature(p), (i, p))
                self.assertEqual(got["status"], want["status"], (i, p))
                self.assertEqual(got["points"], "\n\n".join(want["points"]), (i, p))
                self.assertEqual(got["extra"] or "", "\n".join(want["extra"]), (i, p))
                self.assertEqual(got["dignity"], want["dignity"]["label"], (i, p))
                self.assertEqual(got["dignity_line"] or "", want["dignity_line"], (i, p))
                self.assertEqual(got["digbala_line"] or "", want["digbala_line"], (i, p))
                self.assertEqual(got["classes"] or "", ", ".join(want["classes"]), (i, p))
                self.assertEqual(got["class_texts"] or "", "\n".join(want["class_texts"]), (i, p))

    def test_new_layer_one_lines_equal_python(self):
        mx = self.mx
        new = ("moon_strength", "nature_lines", "twelfth_line", "conditions", "active")
        for i, c, wb in self.each():
            ws = wb[mx.PREFIX + "Predict_Calc"]
            d = br.dasha_now(RULES, c["signs"]["Moon"] * 30 + c["degs"]["Moon"], c["birth"], self.now)
            ctx = br.context(c["lagna"], c["signs"], c["degs"], c["waxing"], c.get("retro"), c.get("gender"),
                             d["age"], d["maha"], d["bhukti"])
            conds = br.chart_conditions(RULES, ctx)
            for k, p in enumerate(br.PLANET_ORDER):
                r = mx.PREDICT_ROW0 + k
                g = br.graha_bhava(RULES, c["lagna"], c["signs"], p, c["degs"][p], c["waxing"])
                got = {name: ws.cell(row=r, column=mx.PREDICT_COLS[name]).value or "" for name in new}
                want = {"moon_strength": g["moon_strength"], "nature_lines": "\n".join(x["text"] for x in g["nature_lines"]),
                        "twelfth_line": g["twelfth_line"],
                        "conditions": "\n".join(x["text"] for x in conds if x["planet"] == p),
                        "active": ", ".join(x["key"] for x in conds if x["planet"] == p and x["active"])}
                self.assertEqual(got, want, (i, p))

    def test_dignity_labels_equal_python(self):
        mx = self.mx
        seen = set()
        for i, c, wb in self.each():
            ws = wb[mx.PREFIX + "Predict_Calc"]
            for k, p in enumerate(br.PLANET_ORDER):
                got = ws.cell(row=mx.PREDICT_ROW0 + k, column=mx.PREDICT_COLS["dignity"]).value
                want = br.dignity(RULES, p, c["signs"][p], c["degs"][p])["label"]
                self.assertEqual(got, want, (i, p, c["signs"][p], c["degs"][p]))
                seen.add(want)
        self.assertTrue({"Deep Exalted", "Exalted", "Deep Debilitated", "Debilitated", "Own (Moolatrikona)", "Own House",
                         "Friend's House", "Enemy's House"} <= seen, seen)

    def test_dasha_dates_equal_python(self):
        mx = self.mx
        for i, c, wb in self.each(24):
            ws = wb[mx.PREFIX + "Dasha_Calc"]
            moon_lon = c["signs"]["Moon"] * 30 + c["degs"]["Moon"]
            d = br.vimshottari(RULES, moon_lon, c["birth"], self.now)
            rows = [(t["lord"], b["lord"], b["start"], b["end"]) for t in d["timeline"] for b in t["bhuktis"]]
            for k, (maha, bhukti, start, end) in enumerate(rows):
                r = mx.DASHA_ROW0 + k
                self.assertEqual((ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value), (maha, bhukti), (i, k))
                for col, want in ((5, start), (6, end)):
                    got = ws.cell(row=r, column=col).value
                    self.assertLess(abs((got - want).total_seconds()), 2, (i, k, col, got, want))
            self.assertEqual((ws.cell(row=mx.DASHA_CUR_ROW, column=2).value, ws.cell(row=mx.DASHA_CUR_ROW + 1, column=2).value),
                             (d["cur_maha"]["lord"], d["cur_bhukti"]["lord"]), i)

    def test_workbook_keeps_the_data_sheets_and_row_labels(self):
        mx = self.mx
        wb = self.books[0]
        for n in ("README", "Classes", "ClassRules", "Badhaka", "DignityEffect", "Digbala", "DashaRoleText",
                  "GrahaInBhava", "GrahaPair", "BhavaLordIn", "GrahaRashi", "Chart", "Classes_Calc", "Roles_Calc",
                  "Dasha_Calc", "Predict_Calc", "Ref_Calc"):
            self.assertIn(mx.PREFIX + n, wb.sheetnames)
        self.assertEqual([wb[mx.PREFIX + "Roles_Calc"].cell(row=2, column=2 + j).value for j in range(6)], br.ROLE_ORDER)
        self.assertEqual([wb[mx.PREFIX + "Classes_Calc"].cell(row=mx.MATRIX_ROW0, column=3 + k).value for k in range(10)], br.CLASS_ORDER)
        self.assertEqual(wb[mx.PREFIX + "GrahaInBhava"].cell(row=1, column=3).value, "Points")
        for k, p in enumerate(br.PLANET_ORDER):
            self.assertEqual(wb[mx.PREFIX + "Chart"].cell(row=mx.PLANET_ROW0 + k, column=1).value, p)


@unittest.skipUnless(openpyxl, "openpyxl is needed for the workbook checks")
class MasterWorkbookInstall(unittest.TestCase):
    """`make_session23_xlsx.py --install` adds the S23_ sheets to the master workbook and touches nothing else."""

    @staticmethod
    def make_master(path):
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Input"
        ws["A1"], ws["B1"], ws["C1"] = "Name", 21, "=B1*2"
        ws.merge_cells("E1:F2")
        ws["E1"] = "merged"
        dv = DataValidation(type="list", formula1='"a,b"')
        ws.add_data_validation(dv)
        dv.add("A3")
        ws.conditional_formatting.add("B1", CellIsRule(operator="greaterThan", formula=["5"], fill=PatternFill("solid", fgColor="FFFF00")))
        ws.column_dimensions["A"].width = 33
        wb.create_sheet("Reference Data")["A1"] = "House"
        wb.save(path)

    def setUp(self):
        import tempfile
        import make_session23_xlsx as mx
        self.mx = mx
        self.tmp = tempfile.TemporaryDirectory()
        self.master = pathlib.Path(self.tmp.name, "Classification_for_Horoscope_Analysis_v7_1.xlsx")
        self.make_master(self.master)
        self.original = self.master.read_bytes()

    def tearDown(self):
        self.tmp.cleanup()

    def test_install_adds_the_sheets_keeps_a_backup_and_leaves_other_sheets_alone(self):
        done = self.mx.install_in_master(self.master)
        self.assertEqual(done["backup"].read_bytes(), self.original)
        self.assertTrue(done["wrote_data"])
        wb = openpyxl.load_workbook(self.master)
        self.assertEqual(wb.sheetnames[:2], ["Input", "Reference Data"])
        for n in ("README", "Chart", "Classes_Calc", "Roles_Calc", "Dasha_Calc", "Predict_Calc", "Ref_Calc", "Classes",
                  "ClassRules", "Badhaka", "DignityEffect", "Digbala", "DashaRoleText", "GrahaInBhava", "GrahaPair",
                  "BhavaLordIn", "GrahaRashi"):
            self.assertIn("S23_" + n, wb.sheetnames)
        old = openpyxl.load_workbook(done["backup"])
        for ws in old.worksheets:
            new = wb[ws.title]
            self.assertEqual([[c.value for c in row] for row in ws.iter_rows()], [[c.value for c in row] for row in new.iter_rows()])
            self.assertEqual(sorted(map(str, ws.merged_cells.ranges)), sorted(map(str, new.merged_cells.ranges)))
            self.assertEqual(len(ws.data_validations.dataValidation), len(new.data_validations.dataValidation))
            self.assertEqual(len(ws.conditional_formatting), len(new.conditional_formatting))
            self.assertEqual({k: v.width for k, v in ws.column_dimensions.items()}, {k: v.width for k, v in new.column_dimensions.items()})

    def test_the_installed_data_sheets_read_back_as_the_committed_rules(self):
        self.mx.install_in_master(self.master)
        got = b23_read(self.master)
        self.assertEqual(got, {k: v for k, v in RULES_JSON().items() if k not in ("reference", "meta")})

    def test_reinstall_keeps_your_edits_unless_you_reset(self):
        self.mx.install_in_master(self.master)
        wb = openpyxl.load_workbook(self.master)
        ws = wb["S23_GrahaInBhava"]
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Moon" and ws.cell(row=r, column=2).value == 1)
        ws.cell(row=row, column=3).value = "My own reading of the Moon in the 1st."
        wb.save(self.master)
        done = self.mx.install_in_master(self.master)
        self.assertFalse(done["wrote_data"])
        self.assertEqual(openpyxl.load_workbook(self.master)["S23_GrahaInBhava"].cell(row=row, column=3).value, "My own reading of the Moon in the 1st.")
        self.assertEqual(b23_read(self.master)["graha_in_bhava"][row - 2]["points"], ["My own reading of the Moon in the 1st."])
        self.mx.install_in_master(self.master, reset_data=True)
        self.assertNotEqual(openpyxl.load_workbook(self.master)["S23_GrahaInBhava"].cell(row=row, column=3).value, "My own reading of the Moon in the 1st.")

    def test_every_install_keeps_its_own_backup_in_the_backups_folder(self):
        first = self.mx.install_in_master(self.master)
        wb = openpyxl.load_workbook(self.master)
        wb["S23_GrahaInBhava"].cell(row=3, column=3).value = "An edit made before --reset-data."
        wb.save(self.master)
        edited = self.master.read_bytes()
        second = self.mx.install_in_master(self.master, reset_data=True)          # same day, seconds later
        self.assertNotEqual(first["backup"], second["backup"])
        for done in (first, second):
            self.assertEqual(done["backup"].parent, self.master.parent / "backups")
            self.assertRegex(done["backup"].name, r"^Classification_for_Horoscope_Analysis_v7_1\.before-session23-\d{4}-\d{2}-\d{2}_\d{6}(-\d+)?\.xlsx$")
        self.assertEqual(first["backup"].read_bytes(), self.original)
        self.assertEqual(second["backup"].read_bytes(), edited)                    # the edit survives in its backup
        import build_session23 as b23
        self.assertEqual(b23.find_master(self.master.parent), self.master)         # backups never look like the master

    def test_sign_inputs_only_accept_a_sign_from_the_list(self):
        self.mx.install_in_master(self.master)
        lists = [d for d in openpyxl.load_workbook(self.master)["S23_Chart"].data_validations.dataValidation if d.type == "list"]
        self.assertEqual(len(lists), 1)
        self.assertTrue(lists[0].showErrorMessage)          # a typed 'Vrishabha' is refused, not silently accepted
        self.assertIn("B3", str(lists[0].sqref))

    def test_degree_inputs_only_accept_a_degree_within_the_sign(self):
        self.mx.install_in_master(self.master)
        dvs = [d for d in openpyxl.load_workbook(self.master)["S23_Chart"].data_validations.dataValidation if d.type == "decimal"]
        self.assertEqual(len(dvs), 1)
        self.assertIn(f"C{self.mx.PLANET_ROW0}:C{self.mx.PLANET_ROW0 + 8}", str(dvs[0].sqref))
        self.assertEqual((float(dvs[0].formula1), float(dvs[0].formula2)), (0.0, 29.999999))
        self.assertEqual(dvs[0].showErrorMessage, True)

    @unittest.skipUnless(SOFFICE, "LibreOffice is needed to recalculate the installed calculator")
    def test_installed_calculator_gives_the_slide_chart_roles(self):
        import tempfile
        mx = self.mx
        mx.install_in_master(self.master)
        with tempfile.TemporaryDirectory() as out:
            _lo_recalc([self.master], out)
            wb = openpyxl.load_workbook(pathlib.Path(out, self.master.name), data_only=True)
        want = br.planet_roles(RULES, 4, {"Moon": 6, "Ketu": 7, "Saturn": 9, "Jupiter": 10, "Mars": 3, "Sun": 2,
                                          "Venus": 2, "Mercury": 1, "Rahu": 1})
        got = {p: wb["S23_Roles_Calc"].cell(row=mx.ROLES_ROW0 + k, column=8).value or "" for k, p in enumerate(br.PLANET_ORDER)}
        self.assertEqual(got, {p: ", ".join(want[p]) for p in br.PLANET_ORDER})
        self.assertEqual(wb["Input"]["C1"].value, 42)                   # the master's own formulas still calculate


def b23_read(path):
    import build_session23 as b23
    return b23.read_workbook(path)


def RULES_JSON():
    return json.loads(JSON_PATH.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
