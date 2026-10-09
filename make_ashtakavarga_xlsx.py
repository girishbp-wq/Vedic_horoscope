#!/usr/bin/env python3
"""Build Ashtakavarga.xlsx — the Excel twin of ashtakavarga.py and the Ashtakavarga tab.

    python3 make_ashtakavarga_xlsx.py [output.xlsx]

Three sheets, all driven by formulas so the workbook recalculates when you change a sign:

  Ashtakavarga  pick the sign of each contributor (dropdown); BAV for the seven planets,
                SAV, house-from-Lagna and checksum cells ("OK" / "CHECK") update live.
  Rules         the Parashara tables as 1/0 flags per house (edit here to test another
                recension; the checksum cells turn to CHECK if a total is off).
  Workings      one block per planet: what each of the eight contributors gives each sign.

The rule table comes from ashtakavarga.RULES, so Python, Excel and index.html agree.
If LibreOffice is installed the saved file is round-tripped through it once so that the
results are also stored as cached values (previews, phones and OneDrive show numbers
without recalculating).
"""
import pathlib
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

import ashtakavarga as av

SAMPLE = {"Sun": "Taurus", "Moon": "Sagittarius", "Mars": "Aquarius", "Mercury": "Aries",
          "Jupiter": "Gemini", "Venus": "Pisces", "Saturn": "Capricorn", "Lagna": "Taurus"}

MAROON, PARCH, GOLD = "6B1D2B", "F6ECD6", "C9A24C"
GREEN, AMBER, RED = "DCEFD5", "FDF0C8", "F6D9D2"
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=MAROON)
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="D8C69A")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# --- layout contract (the tests rely on these) -------------------------------------
INPUT_ROW0 = 5      # Ashtakavarga!A5:C12  contributor / sign name / sign #
BAV_HEAD, BAV_ROW0 = 15, 16
SAV_NAMES, SAV_HOUSE, SAV_ROW = 25, 26, 27   # heading on row 24, clear of the Saturn row (22)
RULES_ROW0 = 5      # Rules!A5:O60
WORK_HEAD, WORK_INDEX, WORK_ROW0, WORK_BLOCK = 3, 4, 6, 12


def rules_row(target, contributor):
    return RULES_ROW0 + av.PLANETS.index(target) * 8 + av.CONTRIBUTORS.index(contributor)


def bav_row(target):
    """Row on Workings holding the BAV total of `target`."""
    return WORK_ROW0 + av.PLANETS.index(target) * WORK_BLOCK + 9


def _head(ws, row, c0, labels):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=c0 + i, value=text)
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, CENTER, BOX


def _body(ws, rows, c0, c1):
    for r in rows:
        for c in range(c0, c1 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            if c > c0:
                cell.alignment = CENTER


def _rules_sheet(ws):
    ws["A1"] = "Ashtakavarga rules — houses (counted from the contributor, 1 = its own sign) that give a bindu"
    ws["A1"].font = Font(bold=True, size=13, color=MAROON)
    ws["A2"] = ("1 = the contributor gives one bindu to the target planet's chart in that house from itself. "
                "Tables follow Brihat Parashara Hora Shastra, ch. 66. Edit a flag to test another school; "
                "the check cells at right turn to CHECK if a planet's total leaves its classical value.")
    ws["A3"] = ("Moon rows: the BPHS text is used (Moon<-Moon has house 9, Moon<-Mars lacks 9, Moon<-Jupiter has 2 "
                "not 12). Most modern books differ: Moon<-Moon remove 9; Moon<-Mars add 9; Moon<-Jupiter replace 2 with 12.")
    ws["A3"].font = Font(italic=True, color="8A7A55")
    _head(ws, 4, 1, ["Target planet", "Contributor"] + [f"House {h}" for h in range(1, 13)] + ["Bindus"])
    for t in av.PLANETS:
        for c in av.CONTRIBUTORS:
            r = rules_row(t, c)
            ws.cell(row=r, column=1, value=t)
            ws.cell(row=r, column=2, value=c)
            for h in range(1, 13):
                ws.cell(row=r, column=2 + h, value=1 if h in av.RULES[t][c] else 0)
            ws.cell(row=r, column=15, value=f"=SUM(C{r}:N{r})")
    last = RULES_ROW0 + 7 * 8 - 1
    _body(ws, range(RULES_ROW0, last + 1), 1, 15)
    ws.conditional_formatting.add(f"C{RULES_ROW0}:N{last}",
                                  CellIsRule(operator="equal", formula=["1"],
                                             fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    # checksum table
    _head(ws, 4, 17, ["Target", "Classical", "In table", "Check"])
    for i, t in enumerate(av.PLANETS):
        r = RULES_ROW0 + i
        ws.cell(row=r, column=17, value=t)
        ws.cell(row=r, column=18, value=av.EXPECTED_TOTALS[t])
        ws.cell(row=r, column=19, value=f"=SUMIF($A${RULES_ROW0}:$A${last},Q{r},$O${RULES_ROW0}:$O${last})")
        ws.cell(row=r, column=20, value=f'=IF(S{r}=R{r},"OK","CHECK")')
    r = RULES_ROW0 + 7
    ws.cell(row=r, column=17, value="All")
    ws.cell(row=r, column=18, value=f"=SUM(R{RULES_ROW0}:R{r - 1})")
    ws.cell(row=r, column=19, value=f"=SUM(S{RULES_ROW0}:S{r - 1})")
    ws.cell(row=r, column=20, value=f'=IF(S{r}=R{r},"OK","CHECK")')
    _body(ws, range(RULES_ROW0, r + 1), 17, 20)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 12
    for c in range(3, 16):
        ws.column_dimensions[col(c)].width = 8
    for c in "QRST":
        ws.column_dimensions[c].width = 11
    ws.freeze_panes = "C5"


def _workings_sheet(ws):
    ws["A1"] = "Workings — what each contributor gives each sign, planet by planet"
    ws["A1"].font = Font(bold=True, size=13, color=MAROON)
    ws.cell(row=WORK_HEAD, column=1, value="Sign").font = Font(bold=True)
    ws.cell(row=WORK_INDEX, column=1, value="Sign # (0 = Aries)").font = Font(bold=True)
    for s, name in enumerate(av.SIGNS):
        h = ws.cell(row=WORK_HEAD, column=3 + s, value=name)
        h.font, h.fill, h.alignment = HEAD_FONT, HEAD_FILL, CENTER
        ws.cell(row=WORK_INDEX, column=3 + s, value=s).alignment = CENTER
    for k, t in enumerate(av.PLANETS):
        r0 = WORK_ROW0 + k * WORK_BLOCK
        title = ws.cell(row=r0, column=1, value=f"{t} Ashtakavarga (classical total {av.EXPECTED_TOTALS[t]})")
        title.font = Font(bold=True, color=MAROON)
        ws.cell(row=r0, column=2, value="From sign #").font = Font(bold=True)
        ws.cell(row=r0, column=15, value="Bindus given").font = Font(bold=True)
        for i, c in enumerate(av.CONTRIBUTORS):
            r = r0 + 1 + i
            rr = rules_row(t, c)
            ws.cell(row=r, column=1, value=c)
            ws.cell(row=r, column=2, value=f"=Ashtakavarga!$C${INPUT_ROW0 + i}")
            for s in range(12):
                cl = col(3 + s)
                ws.cell(row=r, column=3 + s,
                        value=f"=INDEX(Rules!$C${rr}:$N${rr},MOD({cl}${WORK_INDEX}-$B{r},12)+1)")
            ws.cell(row=r, column=15, value=f"=SUM(C{r}:N{r})")
        rb = r0 + 9
        ws.cell(row=rb, column=1, value=f"{t} BAV").font = Font(bold=True)
        for s in range(12):
            cl = col(3 + s)
            c = ws.cell(row=rb, column=3 + s, value=f"=SUM({cl}{r0 + 1}:{cl}{r0 + 8})")
            c.font = Font(bold=True)
        ws.cell(row=rb, column=15, value=f"=SUM(C{rb}:N{rb})").font = Font(bold=True)
        _body(ws, range(r0 + 1, rb + 1), 1, 15)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 12
    for c in range(3, 16):
        ws.column_dimensions[col(c)].width = 11


def _main_sheet(ws, signs):
    ws["A1"] = "Ashtakavarga — Bhinnashtakavarga (BAV) and Sarvashtakavarga (SAV)"
    ws["A1"].font = Font(bold=True, size=14, color=MAROON)
    ws["A2"] = ("Pick each contributor's sidereal sign in column B (replace the sample chart with yours). "
                "Everything below recalculates. Rahu and Ketu take no part.")
    _head(ws, 4, 1, ["Contributor", "Sign (choose)", "Sign #"])
    for i, c in enumerate(av.CONTRIBUTORS):
        r = INPUT_ROW0 + i
        ws.cell(row=r, column=1, value=c).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=signs[c])
        cell.fill = PatternFill("solid", fgColor="FFF7E4")
        ws.cell(row=r, column=3, value=f"=MATCH(B{r},Workings!$C${WORK_HEAD}:$N${WORK_HEAD},0)-1")
    _body(ws, range(INPUT_ROW0, INPUT_ROW0 + 8), 1, 3)
    dv = DataValidation(type="list", formula1=f"=Workings!$C${WORK_HEAD}:$N${WORK_HEAD}", allow_blank=False,
                        showErrorMessage=True, errorTitle="Sign", error="Choose one of the twelve signs.")
    ws.add_data_validation(dv)
    dv.add(f"B{INPUT_ROW0}:B{INPUT_ROW0 + 7}")

    lagna_row = INPUT_ROW0 + 7
    ws.cell(row=BAV_HEAD - 1, column=1, value="Bhinnashtakavarga (BAV) — one row per planet").font = \
        Font(bold=True, color=MAROON)
    _head(ws, BAV_HEAD, 1, ["Planet"] + list(av.SIGNS) + ["Total", "Classical", "Check"])
    for k, t in enumerate(av.PLANETS):
        r = BAV_ROW0 + k
        ws.cell(row=r, column=1, value=t).font = Font(bold=True)
        for s in range(12):
            ws.cell(row=r, column=2 + s, value=f"=Workings!{col(3 + s)}{bav_row(t)}")
        ws.cell(row=r, column=14, value=f"=SUM(B{r}:M{r})").font = Font(bold=True)
        ws.cell(row=r, column=15, value=av.EXPECTED_TOTALS[t])
        ws.cell(row=r, column=16, value=f'=IF(N{r}=O{r},"OK","CHECK")')
    last_bav = BAV_ROW0 + 6
    _body(ws, range(BAV_ROW0, last_bav + 1), 1, 16)

    ws.cell(row=SAV_NAMES - 1, column=1, value="Sarvashtakavarga (SAV) — all seven planets together").font = \
        Font(bold=True, color=MAROON)
    _head(ws, SAV_NAMES, 1, ["Sign"] + list(av.SIGNS) + ["Total", "Classical", "Check"])
    ws.cell(row=SAV_HOUSE, column=1, value="House from Lagna").font = Font(bold=True)
    ws.cell(row=SAV_ROW, column=1, value="SAV bindus").font = Font(bold=True)
    for s in range(12):
        ws.cell(row=SAV_HOUSE, column=2 + s, value=f"=MOD(Workings!{col(3 + s)}${WORK_INDEX}-$C${lagna_row},12)+1")
        c = ws.cell(row=SAV_ROW, column=2 + s, value=f"=SUM({col(2 + s)}{BAV_ROW0}:{col(2 + s)}{last_bav})")
        c.font = Font(bold=True)
    ws.cell(row=SAV_ROW, column=14, value=f"=SUM(B{SAV_ROW}:M{SAV_ROW})").font = Font(bold=True)
    ws.cell(row=SAV_ROW, column=15, value=av.SAV_TOTAL)
    ws.cell(row=SAV_ROW, column=16,
            value=f'=IF(AND(N{SAV_ROW}=O{SAV_ROW},COUNTIF(P{BAV_ROW0}:P{last_bav},"OK")=7),"OK","CHECK")')
    _body(ws, range(SAV_HOUSE, SAV_ROW + 1), 1, 16)

    notes = ["Shading — BAV: 5 or more bindus strong (green), 3-4 average (amber), fewer than 3 weak (red).",
             "Shading — SAV: 30 or more strong, 25-29 average, fewer than 25 weak. The SAV always totals 337 (about 28 per sign).",
             "Bold cell in the BAV table = that planet's own natal sign. Orange header = the Lagna sign.",
             "These are the unreduced (pre-sodhana) figures: no Trikona or Ekadhipatya reduction is applied."]
    for i, n in enumerate(notes):
        ws.cell(row=SAV_ROW + 2 + i, column=1, value=n).font = Font(italic=True, color="8A7A55")

    def fill(c):
        return PatternFill("solid", bgColor=c, fgColor=c)

    rng = f"B{BAV_ROW0}:M{last_bav}"
    # The planet's own natal sign: bold + boxed.  One rule per shading band, and added first,
    # because LibreOffice applies only the first matching condition (Excel merges them).
    first = f"B{BAV_ROW0}"
    natal = f"COLUMN({first})-2=INDEX($C${INPUT_ROW0}:$C${INPUT_ROW0 + 6},ROW({first})-{BAV_ROW0 - 1})"
    med = Side(style="medium", color=MAROON)
    for band, colour in ((f"{first}>=5", GREEN), (f"AND({first}>=3,{first}<=4)", AMBER), (f"{first}<3", RED)):
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"AND({natal},{band})"], fill=fill(colour), font=Font(bold=True, color=MAROON),
            border=Border(left=med, right=med, top=med, bottom=med)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["5"], fill=fill(GREEN)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["3", "4"], fill=fill(AMBER)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["3"], fill=fill(RED)))
    rng = f"B{SAV_ROW}:M{SAV_ROW}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["30"], fill=fill(GREEN)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["25", "29"], fill=fill(AMBER)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["25"], fill=fill(RED)))
    for r in (BAV_HEAD, SAV_NAMES):
        ws.conditional_formatting.add(
            f"B{r}:M{r}", FormulaRule(formula=[f"COLUMN(B{r})-2=$C${lagna_row}"],
                                      fill=PatternFill("solid", bgColor="D97324", fgColor="D97324")))
    for rng in (f"P{BAV_ROW0}:P{last_bav}", f"P{SAV_ROW}"):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CHECK"'],
                                                      font=Font(bold=True, color="FFFFFF"), fill=fill("C0392B")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'],
                                                      font=Font(bold=True, color="2A6F6A")))
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 14
    for c in range(3, 14):
        ws.column_dimensions[col(c)].width = 11
    for c in "NOP":
        ws.column_dimensions[c].width = 10


def build(path, signs=None):
    """Write the workbook (formulas only, no cached values) to `path`."""
    signs = dict(SAMPLE if signs is None else signs)
    wb = Workbook()
    main = wb.active
    main.title = "Ashtakavarga"
    _main_sheet(main, signs)
    _rules_sheet(wb.create_sheet("Rules"))
    _workings_sheet(wb.create_sheet("Workings"))
    for ws in wb.worksheets:  # print landscape, one page wide
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def _store_cached_values(path):
    """Round-trip through LibreOffice so results are saved as cached values too."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("note: LibreOffice not found — workbook saved without cached values "
              "(Excel recalculates it on open).", file=sys.stderr)
        return
    with tempfile.TemporaryDirectory() as tmp:
        out, profile = pathlib.Path(tmp, "out"), pathlib.Path(tmp, "profile").resolve()
        out.mkdir()
        subprocess.run([soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless",
                        "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                       check=True, capture_output=True, timeout=300)
        shutil.copyfile(out / pathlib.Path(path).name, path)


def _console_safe():
    """Never crash on a character the console's code page lacks (Windows cp1252 and the like)."""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    _console_safe()
    if argv and argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0
    path = pathlib.Path(argv[0] if argv else pathlib.Path(__file__).with_name("Ashtakavarga.xlsx"))
    av.validate_rules()
    build(path)
    _store_cached_values(path)
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
