#!/usr/bin/env python3
"""One-time fixes for the master workbook from the October 2026 audit.

    python patch_master_workbook.py [--workbook FILE] [--dry-run]

Close the workbook in Excel first. The script copies it to the backups folder, then changes each
listed cell only if the cell still holds exactly what the audit found there. A cell you have
changed since is left alone and reported as "skipped", so nothing of yours is overwritten. Running
it again is safe: cells already fixed are reported as such. Afterwards open the workbook in Excel
and save once, so the new formulas' results are stored with it.

What it fixes:
  Analysis   Chara/Sthira/Dwisabhava and element counts (B64:B71), the Kuja Dosha checks and
             verdict (B90, A92:C92, B94, D84), the Vysya Moon text (C37), a deep-debilitated Moon
             (B80), "House" with a blank Lagna (B46)
  Input      #VALUE! in the house column with a blank Lagna (D13:D21), deep exaltation and
             debilitation within 1 degree as on the page (G13:G19), Rahu/Ketu outside exaltation
             or debilitation shown as "Node (no rulership)" (G20:G21), the degree header (E11)
  Tithi      tithi 15 is Purnima and 30 Amavasya (C18, C33); the first tithi spans 0-12 deg (C36)
  Data       Kala Purusha body parts for Karkataka-Tula as in the teacher's table on Sheet1
             (Reference Data T5:T8); Mercury's letters are the retroflex Ta-varga; Mercury and
             Venus go round the zodiac in about a year; Bharani's yoni is male and its varna
             Mleccha (Revati is the female elephant); Rohini's purushartha is Moksha; Mercury's
             Moolatrikona starts at 15 deg; Makara (not Vrischika) is Chara; "0-20" Moolatrikona;
             the Shukla/Krishna rule in Planet Characteristics; the Nakshatras lord header.
"""
import pathlib
import sys

from openpyxl import load_workbook

import build_session23 as b23
from make_session23_xlsx import backup_master


class Patch:
    """Set sheet!cell to `new` if it holds `old`. With part=True, `old` -> `new` is a substring
    replacement (every occurrence) inside the cell's text."""

    def __init__(self, sheet, cell, old, new, why, part=False):
        self.sheet, self.cell, self.old, self.new, self.why, self.part = sheet, cell, old, new, why, part


def _text(v):
    return getattr(v, "text", v)          # an array formula's text


P = []

# ---- Analysis: counts by modality and element. INDEX(...) with an array of rows returns only the
#      first item in Excel, so the saved counts were 9/0/0. COUNTIFS on the zodiac numbers works everywhere.
_COUNT_OLD = ("=IF(Input!C13=\"\",\"\",SUMPRODUCT((Input!C13:C21<>\"\")*1*(INDEX('Reference Data'!{col}$2:{col}$13,"
              "IF(Input!C13:C21=\"\",1,Input!C13:C21))=\"{val}\")))")
_COUNT_NEW = ("=IF(Input!C13=\"\",\"\",SUMPRODUCT(COUNTIFS('Reference Data'!$A$2:$A$13,Input!C13:C21,"
              "'Reference Data'!${col}$2:${col}$13,\"{val}\")))")
for row, col, val in ((64, "E", "Chara"), (65, "E", "Sthira"), (66, "E", "Dwisabhava"), (68, "F", "Agni"),
                      (69, "F", "Prithvi"), (70, "F", "Vayu"), (71, "F", "Jala")):
    P.append(Patch("Analysis", f"B{row}", _COUNT_OLD.format(col=col, val=val), _COUNT_NEW.format(col=col, val=val),
                   f"count of planets in {val} signs"))

P.append(Patch("Analysis", "C37", 'IF(B37="Vaishya",', 'IF(OR(B37="Vaishya",B37="Vysya"),',
               "the data spells the varna Vysya", part=True))
P.append(Patch("Analysis", "B80", 'Input!G14="⬇ Debilitated"', 'ISNUMBER(SEARCH("Debilitated",Input!G14))',
               "a deep-debilitated Moon also needs the remedy", part=True))
P.append(Patch("Analysis", "B46", '=IF(Input!C14="","","House "&Input!D14)', '=IF(Input!D14="","","House "&Input!D14)',
               "no half-filled 'House' while the Lagna is blank"))

# ---- Analysis: Kuja Dosha (teacher's rules, Planet Characteristics rows 113-119)
P.append(Patch("Analysis", "D84",
               '=IF(B84="","",IF(OR(B84=7,B84=8),"HIGH (7th/8th from Lagna)",IF(OR(B84=2,B84=4,B84=12),"Moderate","—")))',
               '=IF(B84="","",IF(OR(B84=7,B84=8),"Severe (7th/8th from Lagna)",IF(OR(B84=2,B84=4,B84=12),'
               '"Strong (Lagna gives the strongest impact)","—")))',
               "Lagna is the strongest reference point, as on the page"))
P.append(Patch("Analysis", "B90",
               '=IF(Input!G15="","",IF(OR(ISNUMBER(SEARCH("Exalted",Input!G15)),ISNUMBER(SEARCH("Own House",Input!G15)),'
               'ISNUMBER(SEARCH("Moola",Input!G15))),"✅ Cancelled","❌ Not Met"))',
               '=IF(Input!G15="","",IF(OR(ISNUMBER(SEARCH("Exalted",Input!G15)),ISNUMBER(SEARCH("Own House",Input!G15)),'
               'ISNUMBER(SEARCH("Moola",Input!G15)),Input!C15=5,Input!C15=9,Input!C15=12),"✅ Cancelled","❌ Not Met"))',
               "rule 1 includes a friend's house: Simha, Dhanus, Meena"))
P.append(Patch("Analysis", "C90", '=IF(Input!G15="","","Mars Dignity: "&Input!G15)',
               '=IF(Input!G15="","","Mars Dignity: "&Input!G15&IF(OR(Input!C15=5,Input!C15=9,Input!C15=12),'
               '" — a friend\'s house",""))',
               "say when the friend's-house rule applies"))
_SAT = "MOD(Input!C15-Input!C19,12)+1"                                  # Mars counted from Saturn
_NAK = "MIN(27,INT(((Input!C14-1)*30+Input!E14)*3/40)+1)"               # Moon's nakshatra, 1-27
_EXEMPT = "{1,5,7,8,9,12,15,17,22,26,27}"
P.append(Patch("Analysis", "A92", "3. Other cancellation factors (manual check)",
               "3. Saturn conjoins or aspects Mars, or an exempt birth nakshatra (Mangalik partner: check at matching)",
               "rules 2 and 4 are now worked out"))
P.append(Patch("Analysis", "B92", "ℹ Check manually",
               f'=IF(OR(Input!C15="",Input!C19="",Input!C14="",Input!E14=""),"",IF(OR(ISNUMBER(MATCH({_SAT},{{1,3,7,10}},0)),'
               f'ISNUMBER(MATCH({_NAK},{_EXEMPT},0))),"✅ Cancelled","❌ Not Met"))',
               "Saturn conjunction/aspect and the birth nakshatra"))
P.append(Patch("Analysis", "C92",
               "Also cancelled if: (a) Mars conjoined/aspected by Saturn; (b) Born in Ashwini, Mrigashira, Punarvasu, Pushya, "
               "Ashlesha, Uttara, Swati, Anuradha, Shravana, Uttara Bhadra, or Revathi nakshatra; (c) Marriage between two "
               "Mangalik individuals.",
               f'=IF(OR(Input!C15="",Input!C19="",Input!C14="",Input!E14=""),"Needs the Mars, Saturn and Moon signs and the '
               f'Moon\'s degree.","Saturn in "&Input!B19&": "&IF({_SAT}=1,"conjoins Mars ✅",IF(ISNUMBER(MATCH({_SAT},{{3,7,10}},0)),'
               f'"casts its "&IF({_SAT}=3,"3rd",{_SAT}&"th")&" aspect on Mars ✅","no conjunction or aspect on Mars"))&". Birth '
               f'nakshatra: "&INDEX(Nakshatras!$B$5:$B$31,{_NAK})&IF(ISNUMBER(MATCH({_NAK},{_EXEMPT},0))," — exempt ✅",'
               f'" — not exempt")&". Exempt: Ashwini, Mrigashira, Punarvasu, Pushya, Ashlesha, Uttara (Phalguni), Swati, '
               f'Anuradha, Shravana, Uttara Bhadra, Revathi. A marriage between two Mangaliks also cancels it — check at matching.")',
               "explain rules 2 and 4 for this chart"))
P.append(Patch("Analysis", "B94",
               '=IF(OR(B84="",B85="",B86=""),"",IF(OR(B90="✅ Cancelled",B91="✅ Cancelled"),"✅ Kuja Dosha CANCELLED '
               '(auto-detected). Verify manual cancellation factors as well.",IF(AND(LEFT(C84,1)="—",LEFT(C85,1)="—",'
               'LEFT(C86,1)="—"),"✅ No Kuja Dosha from any reference point.",IF(B93="✅ Reduced","⚠ Kuja Dosha PRESENT but '
               'REDUCED (native over 28). Verify other manual cancellations.","⚠ Kuja Dosha PRESENT. Severity depends on source '
               '(Lagna strongest). Verify manual cancellation factors (Saturn aspect, nakshatra, partner status)."))))',
               '=IF(B84="","",IF(AND(LEFT(C84,1)<>"⚠",LEFT(C85,1)<>"⚠",LEFT(C86,1)<>"⚠"),"✅ No Kuja Dosha from any '
               'reference point.",IF(OR(B90="✅ Cancelled",B91="✅ Cancelled",B92="✅ Cancelled"),"✅ Kuja Dosha present but '
               'CANCELLED (rule "&SUBSTITUTE(TRIM(IF(B90="✅ Cancelled","1 ","")&IF(B91="✅ Cancelled","2 ","")&'
               'IF(B92="✅ Cancelled","3",""))," ",", ")&").",IF(B93="✅ Reduced","⚠ Kuja Dosha PRESENT but REDUCED (native '
               'over 28). At matching, a Mangalik partner also cancels it.","⚠ Kuja Dosha PRESENT and not cancelled by rules '
               '1–3. Severity depends on the reference point (Lagna strongest). At matching, a Mangalik partner cancels it."))))',
               "'no dosha' is decided before any cancellation; rule 3 is now worked out"))

# ---- Input: house numbers with a blank Lagna; deep degrees within 1 deg (the page's tolerance); nodes
for r in range(13, 22):
    P.append(Patch("Input", f"D{r}", f'=IF(C{r}="","",MOD(C{r}-$C$12,12)+1)', f'=IF(OR(C{r}="",$C$12=""),"",MOD(C{r}-$C$12,12)+1)',
                   "no #VALUE! while the Lagna is blank"))
_DEEP = {13: [("ROUND(E13,0)=10", "ABS(E13-10)<=1")],
         14: [("AND(C14=2,ROUND(E14,0)=3)", "AND(C14=2,E14>=2,E14<=3)"), ("AND(C14=8,ROUND(E14,0)=3)", "AND(C14=8,ABS(E14-3)<=1)")],
         15: [("ROUND(E15,0)=28", "ABS(E15-28)<=1")],
         16: [("AND(C16=6,ROUND(E16,0)=15)", "AND(C16=6,E16>=14,E16<=15)"), ("AND(C16=12,ROUND(E16,0)=15)", "AND(C16=12,ABS(E16-15)<=1)")],
         17: [("ROUND(E17,0)=5", "ABS(E17-5)<=1")],
         18: [("ROUND(E18,0)=27", "ABS(E18-27)<=1")],
         19: [("ROUND(E19,0)=20", "ABS(E19-20)<=1")]}
for r, pairs in _DEEP.items():
    for old, new in pairs:
        P.append(Patch("Input", f"G{r}", old, new, "deep exaltation/debilitation within 1 degree, as on the page "
                       "(for the Moon and Mercury only up to the exact degree, where Moolatrikona begins)", part=True))
for r in (20, 21):
    P.append(Patch("Input", f"G{r}", '"⬇ Debilitated","—")))', '"⬇ Debilitated","Node (no rulership)")))',
                   "Rahu and Ketu own no sign", part=True))
P.append(Patch("Input", "E11", "Degree (0-30)", "Degree (decimal 0–30; 11°15′ = 11.25)",
               "degrees are decimals, not degrees.minutes"))

# ---- Tithi
P.append(Patch("Tithi", "C18", "Purnima / Amavasya", "Purnima", "tithi 15 is Purnima"))
P.append(Patch("Tithi", "C33", "Purnima / Amavasya", "Amavasya", "tithi 30 is Amavasya"))
P.append(Patch("Tithi", "C36", "When the Moon is 12 degrees ahead of the Sun, the first tithi begins;",
               "The first tithi (Pratipada) runs while the Moon is 0 to 12 degrees ahead of the Sun;",
               "Pratipada is 0–12 degrees", part=True))

# ---- Data: teacher's table and the workbook's own sheets
for cell, old, new in (("T5", "Chest, Lungs, Breasts, Ribcage, Stomach (upper)", "Heart, Lungs, Chest, Breasts"),
                       ("T6", "Heart, Spine, Upper back, Bones, Eyes (Sun's karaka)", "Stomach, Womb, Upper abdomen"),
                       ("T7", "Stomach (lower), Intestines, Digestive system, Abdomen, Navel", "Hip, Waist, Intestines"),
                       ("T8", "Lower abdomen, Kidneys, Lumbar region, Lower back, Skin", "Kidneys, Lower abdomen, Lumbar region")):
    P.append(Patch("Reference Data", cell, old, new, "Kala Purusha body part as in the teacher's table (Sheet1)"))
_TA_OLD, _TA_NEW = "त थ द ध न (Ta, Tha, Da, Dha, Na)", "ट ठ ड ढ ण (Ṭa, Ṭha, Ḍa, Ḍha, Ṇa)"
P.append(Patch("Planet Characteristics", "E60", _TA_OLD, _TA_NEW, "Mercury's letters are the Ta-varga (Jupiter has the ta-varga)"))
P.append(Patch("Planet Profile", "AS8", _TA_OLD, _TA_NEW, "Mercury's letters are the Ta-varga (Jupiter has the ta-varga)"))
for sheet, cell in (("Planet Characteristics", "E14"), ("Planet Profile", "G8")):
    P.append(Patch(sheet, cell, "~288 days", "~1 year (moves with the Sun)", "Mercury goes round the zodiac in about a year"))
for sheet, cell in (("Planet Characteristics", "G14"), ("Planet Profile", "G10")):
    P.append(Patch(sheet, cell, "~324 days", "~1 year (moves with the Sun)", "Venus goes round the zodiac in about a year"))
P.append(Patch("Planet Characteristics", "C84",
               "Count houses from Sun's rashi to Moon's rashi. If Moon falls in houses 1–7 from Sun → Shukla Paksha (waxing). "
               "If Moon falls in houses 8–12 from Sun → Krishna Paksha (waning).",
               "Take the Moon's distance ahead of the Sun: 0°–180° → Shukla Paksha (waxing), 180°–360° → Krishna Paksha "
               "(waning) — the tithi (Tithi sheet). Counting signs from the Sun's rashi to the Moon's is only a rough guide "
               "and is wrong near the edges: Sun 20° Mesha with Moon 5° Mesha is Krishna Chaturdashi.",
               "paksha from the Sun–Moon distance"))
P.append(Patch("Nakshatra Detail", "H6", "Female", "Male", "Bharani is the male elephant (Revati is the female)"))
P.append(Patch("Nakshatra Detail", "D6", "Vaishya", "Mleccha", "Bharani's varna"))
P.append(Patch("Nakshatra Detail", "C8", None, "Moksha", "Rohini's purushartha"))
P.append(Patch("Planet Dignity", "I7", "16°–20° Kanya (Mercury)", "15°–20° Kanya (Mercury)",
               "Mercury's Moolatrikona is 15°–20° (Planet Characteristics, Planets)"))
P.append(Patch("Explanations", "B15", "Tula , Vrischika )", "Tula , Makara )", "Makara is the fourth Chara sign", part=True))
P.append(Patch("Guide & Explanations", "B208", "(since 0–0 is Moola Trikona)", "(since 0–20 is Moola Trikona)",
               "Simha's Moolatrikona is 0–20 degrees", part=True))
P.append(Patch("Nakshatras", "C4", "Pada lord", "Nakshatra lord", "the column holds the nakshatra's lord"))

PATCHES = P


def apply_patches(wb, patches=PATCHES):
    """Apply to an open workbook; returns [(status, patch, current_text)] with status in
    'fixed', 'already fixed', 'skipped' (the cell holds something else), 'no sheet'."""
    out = []
    for p in patches:
        if p.sheet not in wb.sheetnames:
            out.append(("no sheet", p, None))
            continue
        cell = wb[p.sheet][p.cell]
        cur = _text(cell.value)
        if p.part:
            text = "" if cur is None else str(cur)
            if p.old in text:
                cell.value = text.replace(p.old, p.new)
                out.append(("fixed", p, cur))
            elif p.new in text:
                out.append(("already fixed", p, cur))
            else:
                out.append(("skipped", p, cur))
        elif cur == p.old or (p.old is None and cur in (None, "")):
            cell.value = p.new
            out.append(("fixed", p, cur))
        elif cur == p.new:
            out.append(("already fixed", p, cur))
        else:
            out.append(("skipped", p, cur))
    return out


def patch_master(path, dry_run=False):
    """Back up and patch the workbook at `path`; returns (backup or None, results)."""
    path = pathlib.Path(path)
    wb = load_workbook(path)
    results = apply_patches(wb)
    backup = None
    if not dry_run and any(s == "fixed" for s, _, _ in results):
        backup = backup_master(path, "before-audit-fixes")
        wb.save(path)
    return backup, results


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    for stream in (sys.stdout, sys.stderr):      # never crash on a character the console's code page lacks
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass
    if "-h" in argv or "--help" in argv:
        print(__doc__)
        return 0
    path = argv[argv.index("--workbook") + 1] if "--workbook" in argv else b23.find_master()
    if not path or not pathlib.Path(path).exists():
        raise SystemExit("No Classification_for_Horoscope_Analysis_*.xlsx workbook found next to this script.")
    dry = "--dry-run" in argv
    backup, results = patch_master(path, dry_run=dry)
    counts = {}
    for status, p, cur in results:
        counts[status] = counts.get(status, 0) + 1
        line = f"{status:13s} {p.sheet}!{p.cell} — {p.why}"
        if status == "skipped":
            line += f"\n              (left as it is: the cell no longer holds what the audit found — now {str(cur)[:80]!r})"
        print(line)
    print()
    print(", ".join(f"{n} {s}" for s, n in counts.items()))
    if dry:
        print("Dry run: nothing was saved.")
    elif backup:
        print(f"Backup: {backup}")
        print("Open the workbook in Excel and save once, so the new formulas' results are stored with it.")
    else:
        print("Nothing to change.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
